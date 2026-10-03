import asyncio
from datetime import datetime
import os
from pathlib import Path
import sqlite3
import sys
import tempfile
import unittest

TEST_DIR = tempfile.TemporaryDirectory(prefix="firm-security-")
os.environ["FIRM_MANAGER_DATA_DIR"] = TEST_DIR.name
os.environ['FIRM_MANAGER_DEVELOPMENT'] = '1'
os.environ["FIRM_MANAGER_ALLOWED_HOSTS"] = "testserver"
os.environ["FIRM_MANAGER_ALLOWED_PORTS"] = "80"
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))
import httpx
import main
import models
from auth import get_password_hash
from database import SessionLocal, engine
from schema_init import initialize_schema
from sqlalchemy import create_engine


class SecurityTests(unittest.IsolatedAsyncioTestCase):
    @classmethod
    def setUpClass(cls):
        from data_crypto import initialize_key
        initialize_key()
        from license_manager import manager
        manager.identity(initialize=True)
        with SessionLocal() as db:
            password = get_password_hash("Test-password-2026")
            admin = models.User(username="admin", real_name="管理", role_records=[models.UserRoleRecord(role_code="office_admin")], fiscal_year=2026, hashed_password=password)
            leader = models.User(username="leader", real_name="负责人", is_practitioner=True, role_records=[models.UserRoleRecord(role_code="practitioner")], fiscal_year=2026, hashed_password=password)
            outsider = models.User(username="outsider", real_name="无关人员", is_practitioner=True, role_records=[models.UserRoleRecord(role_code="practitioner")], fiscal_year=2026, hashed_password=password)
            db.add_all([admin, leader, outsider, models.FiscalYear(year=2026)])
            db.flush()
            cls.admin_id = admin.id
            cls.leader_id = leader.id
            cls.outsider_id = outsider.id
            for name, roles, practitioner in [("number", ["number_manager"], False), ("finance", ["finance"], False), ("clerk", ["clerk"], False), ("combined", ["number_manager", "practitioner"], True)]:
                db.add(models.User(username=name, real_name=name, fiscal_year=2026, hashed_password=password,
                       is_practitioner=practitioner, role_records=[models.UserRoleRecord(role_code=r) for r in roles]))
            customer = models.Customer(name="测试客户", tax_id="TEST-TAX", identity_status="confirmed")
            db.add(customer)
            db.flush()
            cls.customer_id = customer.id
            db.add(models.Project(project_id="PRJ-2026-0001", firm="测试所", report_type="审计", report_year=2026,
                                  customer_name="测试客户", customer_id=customer.id, contract_amount=1234.56,
                                  leader_id=leader.id, fiscal_year=2026))
            for year in (2025, 2026):
                fy = db.query(models.FiscalYear).filter_by(year=year).first()
                if not fy:
                    fy = models.FiscalYear(year=year); db.add(fy); db.flush()
                firm = models.FiscalYearFirm(fiscal_year_id=fy.id, firm="测试所")
                db.add(firm); db.flush()
                rule = models.ReportNumberRule(fiscal_year_firm_id=firm.id, rule_name="审计编号", template="测试〔{yyyy}〕{nnnn}", sequence_digits=4)
                db.add(rule); db.flush()
                db.add(models.FiscalYearReportType(fiscal_year_firm_id=firm.id, report_type="审计", rule_id=rule.id))
            db.commit()

    @classmethod
    def tearDownClass(cls):
        from runtime_log import logger
        for handler in logger.handlers[:]:
            handler.close();logger.removeHandler(handler)
        engine.dispose()
        TEST_DIR.cleanup()

    async def asyncSetUp(self):
        self.client = httpx.AsyncClient(transport=httpx.ASGITransport(app=main.app), base_url="http://testserver")

    async def asyncTearDown(self):
        await self.client.aclose()

    async def login(self, username="admin"):
        response = await self.client.post("/api/auth/login", json={"username": username, "password": "Test-password-2026"})
        self.assertEqual(response.status_code, 200, response.text)
        self.client.headers["X-CSRF-Token"] = self.client.cookies["firm_csrf"]
        return response

    async def test_anonymous_data_and_activation_denied(self):
        self.assertEqual((await self.client.get("/api/projects")).status_code, 401)
        self.assertEqual((await self.client.post("/api/license/activate", json={"license_document": {}})).status_code, 401)

    async def test_readonly_allows_account_protection_without_privilege_change(self):
        from unittest.mock import patch
        await self.login()
        password=get_password_hash('Test-password-2026')
        with SessionLocal() as db:
            user=models.User(username='readonly-victim',real_name='安全测试',fiscal_year=2026,hashed_password=password,role_records=[models.UserRoleRecord(role_code='clerk')])
            db.add(user);db.commit();uid=user.id;revision=user.permission_revision
        with patch.object(main,'license_required',return_value=True), patch.object(main,'license_status',return_value={'allowed':False,'mode':'expired_readonly','reason':'只读'}):
            response=await self.client.put('/api/users/'+str(uid),json={'expected_revision':revision,'is_active':False,'roles':['office_admin']})
            self.assertEqual(response.status_code,402,response.text)
            response=await self.client.put('/api/users/'+str(uid),json={'expected_revision':revision,'is_active':False})
            self.assertEqual(response.status_code,200,response.text)
            self.assertEqual((await self.client.post('/api/projects',json={})).status_code,402)

    async def test_unrelated_get_and_empty_update_denied(self):
        await self.login("outsider")
        for method in ("GET", "PUT"):
            response = await self.client.request(method, "/api/projects/1", **({"json": {}} if method == "PUT" else {}))
            self.assertIn(response.status_code, (403, 404), response.text)

    async def test_authorized_empty_update_rejected(self):
        await self.login("leader")
        response = await self.client.put("/api/projects/1", json={})
        self.assertEqual(response.status_code, 422)
        self.assertEqual(response.json()["detail"]["code"], "empty_update")

    async def test_unknown_fields_rejected(self):
        await self.login("leader")
        response = await self.client.put("/api/projects/1", json={"report_no": "伪造编号"})
        self.assertEqual(response.status_code, 422)

    async def new_project(self, year=2026, **changes):
        payload = {"fiscal_year": year, "firm": "测试所", "report_type": "审计", "report_year": year,
                   "customer_name": "测试客户", "customer_id": self.customer_id, "leader_id": self.leader_id}
        payload.update(changes)
        result = await self.client.post("/api/projects", json=payload)
        self.assertEqual(result.status_code, 200, result.text)
        return result.json()

    async def test_five_roles_projection_and_finance_boundary(self):
        await self.login("number")
        project = (await self.client.get("/api/projects/1")).json()
        self.assertIsNone(project["contract_amount"])
        self.assertNotIn("phone", project["leader"])
        self.assertNotIn("roles", project["leader"])
        self.assertEqual((await self.client.get("/api/projects/1/finance/invoices")).status_code, 403)
        self.assertEqual((await self.client.patch("/api/projects/1", json={"expected_revision": project["revision"], "contract_amount": 1})).status_code, 403)
        await self.login("finance")
        self.assertEqual((await self.client.get("/api/projects/1")).json()["contract_amount"], 1234.56)
        self.assertEqual((await self.client.post("/api/numbered-years", json={"year": 2027})).status_code, 403)
        self.assertEqual((await self.client.get("/api/users")).status_code, 403)
        await self.login("combined")
        profile = (await self.client.get("/api/auth/me")).json()
        self.assertTrue(profile["is_practitioner"])
        self.assertIn("number.configure", profile["permissions"])
        self.assertNotIn("finance.write", profile["permissions"])
        with SessionLocal() as db:
            self.assertEqual(db.query(models.User).filter_by(is_active=True, is_practitioner=True).count(), 3)

    async def test_member_money_hidden_and_injection_rejected(self):
        await self.login()
        project = await self.new_project(member_ids=[self.outsider_id], contract_amount=123)
        await self.login("outsider")
        result = await self.client.get(f'/api/projects/{project["id"]}')
        self.assertEqual(result.status_code, 200)
        self.assertIsNone(result.json()["contract_amount"])
        self.assertIsNone(result.json()["invoiced_amount"])
        for payload in ({"contract_amount": 9}, {"leader_id": self.outsider_id}, {"fiscal_year": 2025}, {"report_no": "BAD"}):
            result = await self.client.patch(f'/api/projects/{project["id"]}', json={"expected_revision": 1, **payload})
            self.assertIn(result.status_code, (403, 422))

    async def test_per_tab_year_and_revision_conflict(self):
        await self.login()
        first = await self.new_project(2025)
        await self.client.put("/api/users/current-fiscal-year", json={"fiscal_year": 2026})
        second = await self.new_project(2025)
        self.assertEqual(first["fiscal_year"], second["fiscal_year"])
        result = await self.client.get("/api/projects", params={"fiscal_year": 2025})
        self.assertTrue(all(p["fiscal_year"] == 2025 for p in result.json()["items"]))
        url = f'/api/projects/{first["id"]}'
        self.assertEqual((await self.client.patch(url, json={"expected_revision": 1, "priority": "高"})).status_code, 200)
        conflict = await self.client.patch(url, json={"expected_revision": 1, "priority": "低"})
        self.assertEqual(conflict.status_code, 409)
        self.assertEqual((await self.client.get(url)).json()["priority"], "高")

    async def test_number_idempotency_void_and_immutable_customer(self):
        await self.login()
        project = await self.new_project()
        url = f'/api/projects/{project["id"]}'
        headers = {"Idempotency-Key": f'test-number-{project["id"]}'}
        first = await self.client.post(url + "/generate-report-no", json={"expected_revision": 1}, headers=headers)
        self.assertEqual(first.status_code, 200, first.text)
        retry = await self.client.post(url + "/generate-report-no", json={"expected_revision": 1}, headers=headers)
        self.assertEqual(first.json(), retry.json())
        wrong = await self.client.post(url + "/generate-report-no", json={"expected_revision": 2}, headers=headers)
        self.assertEqual(wrong.status_code, 409)
        changed = await self.client.patch(url, json={"expected_revision": 2, "customer_name": "伪造"})
        self.assertEqual(changed.status_code, 409)
        self.assertEqual((await self.client.delete(url, params={"expected_revision": 2})).status_code, 403)
        void = await self.client.post(url + "/void-report-no", json={"expected_revision": 2, "reason": "验收作废"})
        self.assertEqual(void.status_code, 200, void.text)
        self.assertEqual((await self.client.post(url + "/generate-report-no", json={"expected_revision": 1}, headers=headers)).status_code, 409)
        with SessionLocal() as db:
            self.assertEqual(db.query(models.ReportNumberHistory).filter_by(project_id=project["id"]).count(), 1)

    async def test_customer_directory_null_tax_and_snapshots(self):
        await self.login()
        data = {"name": "个人验收主体", "type": "individual", "tax_id": None}
        first = await self.client.post("/api/customers", json=data)
        second = await self.client.post("/api/customers", json=data)
        self.assertNotEqual(first.json()["id"], second.json()["id"])
        customer_id = first.json()["id"]
        create_tax = {"name": "并发税号客户", "tax_id": "REPEAT-TAX"}
        outcomes = await asyncio.gather(*(self.client.post("/api/customers", json=create_tax) for _ in range(5)))
        self.assertEqual(sum(r.status_code == 200 for r in outcomes), 1)
        self.assertTrue(all(r.status_code in (200, 409) for r in outcomes))
        project = await self.new_project(customer_id=customer_id)
        issue = await self.client.post(f'/api/projects/{project["id"]}/generate-report-no', json={"expected_revision": 1}, headers={"Idempotency-Key": "pending-identity"})
        self.assertEqual(issue.status_code, 409)
        renamed = await self.client.patch(f"/api/customers/{customer_id}", json={"expected_revision": 1, "name": "更名个人主体"})
        self.assertEqual(renamed.status_code, 200, renamed.text)
        self.assertEqual((await self.client.get(f'/api/projects/{project["id"]}')).json()["customer_name"], data["name"])
        await self.login("clerk")
        lookup = (await self.client.get("/api/customers/lookup", params={"q": "个人"})).json()
        self.assertNotIn("total", lookup)
        self.assertNotIn("tax_id", lookup["items"][0])
        self.assertEqual((await self.client.get(f"/api/customers/{customer_id}")).status_code, 404)

    async def test_last_admin_and_identity_validation(self):
        await self.login()
        revision = (await self.client.get("/api/auth/me")).json()["permission_revision"]
        result = await self.client.put(f"/api/users/{self.admin_id}", json={"expected_revision": revision, "roles": ["clerk"]})
        self.assertEqual(result.status_code, 400, result.text)
        result = await self.client.post("/api/users", json={"username": "invalid-person", "password": "Test-password-2026", "real_name": "错误身份", "roles": ["practitioner"], "is_practitioner": False})
        self.assertEqual(result.status_code, 422)

    async def test_fifty_parallel_number_requests(self):
        await self.login()
        projects = [await self.new_project() for _ in range(50)]
        results = await asyncio.gather(*(self.client.post(f'/api/projects/{p["id"]}/generate-report-no',
            json={"expected_revision": 1}, headers={"Idempotency-Key": f'parallel-number-{p["id"]}'}) for p in projects))
        self.assertTrue(all(r.status_code == 200 for r in results), [r.text for r in results if r.status_code != 200])
        self.assertEqual(len({r.json()["report_no"] for r in results}), 50)
        with SessionLocal() as db:
            self.assertEqual(db.query(models.ReportNumberHistory).filter(models.ReportNumberHistory.project_id.in_([p["id"] for p in projects])).count(), 50)

    async def test_concurrent_last_admin_disable(self):
        await self.login()
        with SessionLocal() as db:
            extra = models.User(username="race_admin", real_name="并发管理员", fiscal_year=2026,
                hashed_password=get_password_hash("Test-password-2026"), role_records=[models.UserRoleRecord(role_code="office_admin")])
            db.add(extra); db.commit(); extra_id = extra.id
        other = httpx.AsyncClient(transport=httpx.ASGITransport(app=main.app), base_url="http://testserver")
        try:
            await other.post("/api/auth/login", json={"username": "race_admin", "password": "Test-password-2026"})
            other.headers["X-CSRF-Token"] = other.cookies["firm_csrf"]
            results = await asyncio.gather(
                self.client.put(f"/api/users/{extra_id}", json={"expected_revision": 1, "is_active": False}),
                other.put(f"/api/users/{self.admin_id}", json={"expected_revision": 1, "is_active": False}))
            self.assertEqual(sum(r.status_code == 200 for r in results), 1, [r.text for r in results])
            with SessionLocal() as db:
                self.assertEqual(db.query(models.User).filter(models.User.is_active == True,
                    models.User.role_records.any(models.UserRoleRecord.role_code == "office_admin")).count(), 1)
        finally:
            await other.aclose()
            with SessionLocal() as db:
                db.get(models.User, self.admin_id).is_active = True
                db.get(models.User, extra_id).is_active = False
                db.commit()

    async def test_ten_thousand_customer_cursor_search(self):
        with SessionLocal() as db:
            db.bulk_insert_mappings(models.Customer, [{"name": f"分页验收客户{i:05d}", "tax_id": f"PAGE-{i:05d}"} for i in range(10000)])
            db.commit()
        await self.login("clerk")
        cursor = 0
        found = []
        for _ in range(5):
            response = await self.client.get("/api/customers/lookup", params={"q": "分页验收", "cursor": cursor, "limit": 50})
            self.assertEqual(response.status_code, 200, response.text)
            data = response.json()
            self.assertEqual(len(data["items"]), 50)
            self.assertNotIn("total", data)
            found.extend(c["id"] for c in data["items"])
            cursor = data["next_cursor"]
        self.assertEqual(len(set(found)), 250)

    async def verified_profile(self, bank='0012345678901234', customer_id=None):
        result = await self.client.post(f'/api/customers/{customer_id or self.customer_id}/billing-profiles', json={
            'label': '验收档案', 'fields': {'bank_account': bank, 'bank_name': '测试银行', 'recipient_email': 'finance@example.test'}})
        self.assertEqual(result.status_code, 200, result.text)
        profile = result.json(); version = profile['versions'][0]
        result = await self.client.post(f'/api/billing-versions/{version["id"]}/verify', json={
            'expected_revision': version['revision'], 'expected_profile_revision': profile['revision'],
            'decision': 'approve', 'verification_note': '财务明确核对主体和银行资料'})
        self.assertEqual(result.status_code, 200, result.text)
        profile['revision'] = result.json()['profile_revision']
        profile['current_verified_version_id'] = version['id']
        return profile, result.json()['version']

    async def test_billing_encryption_version_and_scope(self):
        await self.login()
        profile, version = await self.verified_profile()
        self.assertNotIn('0012345678901234', str(profile))
        self.assertEqual(version['bank_last4'], '1234')
        forbidden = await self.client.patch(f'/api/billing-versions/{version["id"]}', json={'expected_revision': version['revision'], 'fields': {'bank_account': '9999999999999999'}})
        self.assertEqual(forbidden.status_code, 409)
        revealed = await self.client.post(f'/api/billing-versions/{version["id"]}/reveal', json={'purpose': '开票前核对复制'})
        self.assertEqual(revealed.json()['fields']['bank_account'], '0012345678901234')
        self.assertEqual(revealed.headers['Cache-Control'], 'no-store')
        with SessionLocal() as db:
            stored = db.get(models.BillingVersion, version['id'])
            self.assertNotIn('0012345678901234', stored.sensitive_ciphertext)
            self.assertNotIn('finance@example.test', stored.public_json)
            audit = db.query(models.AuditEvent).filter_by(action='billing.reveal', target_id=version['id']).one()
            self.assertTrue(audit.request_id)
        await self.login('outsider')
        self.assertEqual((await self.client.get(f'/api/billing-versions/{version["id"]}')).status_code, 404)
        self.assertEqual((await self.client.post(f'/api/billing-versions/{version["id"]}/reveal', json={'purpose': '猜测编号'})).status_code, 404)

    async def test_invoice_idempotency_snapshot_and_void(self):
        await self.login()
        project = await self.new_project(contract_amount=1000)
        profile, version = await self.verified_profile()
        invoice_url = f'/api/projects/{project["id"]}/finance/invoices'
        data = {'amount': '123.45', 'occurred_on': '2026-10-02', 'expected_project_revision': project['revision'],
            'billing_profile_id': profile['id'], 'billing_version_id': version['id'], 'expected_profile_revision': profile['revision']}
        headers = {'Idempotency-Key': f'invoice-test-{project["id"]}'}
        results = await asyncio.gather(*(self.client.post(invoice_url, json=data, headers=headers) for _ in range(5)))
        self.assertTrue(all(r.status_code == 200 for r in results), [r.text for r in results])
        self.assertEqual(len({r.json()['id'] for r in results}), 1)
        entry = results[0].json()
        snapshot_url = invoice_url + f'/{entry["id"]}/billing-snapshot'
        first = (await self.client.get(snapshot_url)).json()
        self.assertEqual(first['fields']['bank_account'], '0012345678901234')
        change = await self.client.post(f'/api/billing-profiles/{profile["id"]}/versions', json={'expected_profile_revision': profile['revision'], 'fields': {'bank_account': '0099999999999999'}})
        self.assertEqual(change.status_code, 200, change.text)
        newer = change.json()
        verified = await self.client.post(f'/api/billing-versions/{newer["id"]}/verify', json={'expected_revision': newer['revision'], 'expected_profile_revision': profile['revision']+1, 'decision': 'approve', 'verification_note': '银行已变更再次核对'})
        self.assertEqual(verified.status_code, 200, verified.text)
        self.assertEqual((await self.client.get(snapshot_url)).json()['fields']['bank_account'], first['fields']['bank_account'])
        stale = await self.client.post(invoice_url, json={**data, 'expected_project_revision': 2}, headers={'Idempotency-Key': 'invoice-stale-profile'})
        self.assertEqual(stale.status_code, 409)
        void = await self.client.post(invoice_url+f'/{entry["id"]}/void', json={'expected_revision': entry['revision'], 'reason': '登记纠错；另行处理税务系统凭证'})
        self.assertEqual(void.status_code, 200, void.text)
        current = (await self.client.get(f'/api/projects/{project["id"]}')).json()
        self.assertEqual(current['invoiced_amount'], 0)
        self.assertEqual((await self.client.get(snapshot_url)).json()['fields']['bank_account'], first['fields']['bank_account'])
        with SessionLocal() as db:
            self.assertEqual(db.query(models.FinancialEntry).filter_by(project_id=project['id']).count(), 1)
            self.assertEqual(db.query(models.InvoiceSnapshot).filter_by(financial_entry_id=entry['id']).count(), 1)
            self.assertEqual(db.get(models.FinancialEntry, entry['id']).amount_cents, 12345)
            self.assertEqual(db.query(models.AuditEvent).filter_by(action='finance.invoice.create', target_id=entry['id']).count(), 1)
        other = await self.new_project()
        self.assertEqual((await self.client.get(f'/api/projects/{other["id"]}/finance/invoices/{entry["id"]}/billing-snapshot')).status_code, 404)
        await self.login('outsider')
        self.assertEqual((await self.client.get(snapshot_url)).status_code, 404)

    async def test_billing_key_loss_and_ciphertext_authentication(self):
        from data_crypto import KEY_PATH, decrypt, encrypt
        from fastapi import HTTPException
        encrypted, key_id = encrypt({'bank_account': '001234567890'}, 'entity:1')
        with self.assertRaises(HTTPException):
            decrypt(encrypted, key_id, 'entity:2')
        saved = KEY_PATH.with_suffix('.backup')
        KEY_PATH.rename(saved)
        try:
            with self.assertRaises(HTTPException):
                decrypt(encrypted, key_id, 'entity:1')
            self.assertFalse(KEY_PATH.exists())
        finally:
            saved.rename(KEY_PATH)
        self.assertEqual(decrypt(encrypted, key_id, 'entity:1')['bank_account'], '001234567890')

    async def test_leader_submission_finance_verification_and_stale_version(self):
        await self.login('leader')
        profile = (await self.client.post(f'/api/customers/{self.customer_id}/billing-profiles', json={'label':'负责人提交', 'fields':{'bank_account':'000123456789'}})).json()
        version = profile['versions'][0]
        self.assertEqual((await self.client.post(f'/api/billing-versions/{version["id"]}/verify', json={'expected_revision':1,'expected_profile_revision':1,'decision':'approve','verification_note':'越权'})).status_code,403)
        submitted = await self.client.post(f'/api/billing-versions/{version["id"]}/submit', json={'expected_revision':version['revision']})
        self.assertEqual(submitted.status_code,200,submitted.text)
        stale = await self.client.patch(f'/api/billing-versions/{version["id"]}',json={'expected_revision':1,'fields':{'bank_name':'新银行'}})
        self.assertEqual(stale.status_code,409)
        await self.login('finance')
        preview = await self.client.post(f'/api/billing-versions/{version["id"]}/verification-preview',json={'purpose':'核对负责人提交'})
        self.assertEqual(preview.json()['fields']['bank_account'],'000123456789')
        approved = await self.client.post(f'/api/billing-versions/{version["id"]}/verify',json={'expected_revision':submitted.json()['revision'],'expected_profile_revision':profile['revision'],'decision':'approve','verification_note':'已对照开户证明'})
        self.assertEqual(approved.status_code,200,approved.text)
        with SessionLocal() as db:
            stored = db.get(models.BillingVersion,version['id'])
            self.assertIsNotNone(stored.verified_at)
            self.assertEqual(stored.verification_note,'已对照开户证明')

    async def test_clerk_can_manage_only_own_unconfirmed_drafts(self):
        await self.login('clerk')
        proposal = await self.client.post('/api/customer-change-requests',json={'kind':'create','proposal':{'name':'行政自有客户','tax_id':'CLERK-OWN','type':'enterprise'},'reason':'新客户资料'})
        await self.login('finance')
        reviewed = await self.client.post(f'/api/customer-change-requests/{proposal.json()["id"]}/review',json={'expected_revision':1,'decision':'approve','reason':'核对主体证明'})
        self.assertEqual(reviewed.status_code,200,reviewed.text)
        customer_id=reviewed.json()['customer_id']
        await self.login('clerk')
        profile = await self.client.post(f'/api/customers/{customer_id}/billing-profiles',json={'label':'行政草稿','fields':{'bank_account':'000111222333'}})
        self.assertEqual(profile.status_code,200,profile.text)
        version=profile.json()['versions'][0]
        self.assertEqual((await self.client.post(f'/api/billing-versions/{version["id"]}/reveal',json={'purpose':'编辑自有草稿'})).status_code,200)
        self.assertEqual((await self.client.post(f'/api/billing-versions/{version["id"]}/submit',json={'expected_revision':1})).status_code,200)
        self.assertEqual(len((await self.client.get(f'/api/customers/{customer_id}/billing-profiles')).json()),1)
        self.assertEqual((await self.client.get(f'/api/customers/{self.customer_id}/billing-profiles')).status_code,404)
        await self.login('finance')
        self.assertEqual((await self.client.post(f'/api/billing-versions/{version["id"]}/verify',json={'expected_revision':2,'expected_profile_revision':1,'decision':'approve','verification_note':'核对完成'})).status_code,200)
        await self.login('clerk')
        self.assertEqual((await self.client.get(f'/api/customers/{customer_id}/billing-profiles')).json(),[])
        self.assertEqual((await self.client.post(f'/api/billing-versions/{version["id"]}/reveal',json={'purpose':'读取已确认资料'})).status_code,404)

    async def test_merge_preserves_snapshot_and_transfer_removes_old_access(self):
        with SessionLocal() as db:
            source=models.Customer(name='合并原主体',tax_id='MERGE-OLD',identity_status='confirmed')
            target=models.Customer(name='合并保留主体',tax_id='MERGE-NEW',identity_status='confirmed')
            db.add_all([source,target]);db.commit()
            source_id,target_id=source.id,target.id
        await self.login()
        project=await self.new_project(customer_id=source_id,customer_name='合并原主体')
        profile,version=await self.verified_profile(customer_id=source_id)
        url=f'/api/projects/{project["id"]}/finance/invoices'
        result=await self.client.post(url,json={'amount':'9.99','occurred_on':'2026-10-02','expected_project_revision':project['revision'],'billing_profile_id':profile['id'],'billing_version_id':version['id'],'expected_profile_revision':profile['revision']},headers={'Idempotency-Key':'merge-invoice'})
        self.assertEqual(result.status_code,200,result.text)
        snapshot_url=url+f'/{result.json()["id"]}/billing-snapshot'
        merge_data={'target_id':target_id,'expected_revision':1,'expected_target_revision':1,'reason':'同一主体重复建档核对','identity_checked':True,'dry_run':True}
        self.assertEqual((await self.client.post(f'/api/customers/{source_id}/merge',json=merge_data)).json()['project_count'],1)
        self.assertEqual((await self.client.post(f'/api/customers/{source_id}/merge',json={**merge_data,'dry_run':False})).status_code,200)
        latest=(await self.client.get(f'/api/projects/{project["id"]}')).json()
        self.assertEqual(latest['customer_id'],target_id)
        self.assertEqual(latest['customer_name'],'合并原主体')
        await self.login('leader')
        self.assertEqual((await self.client.get(snapshot_url)).json()['fields']['title'],'合并原主体')
        self.assertEqual((await self.client.post(f'/api/billing-versions/{version["id"]}/reveal',json={'purpose':'查看历史'})).status_code,200)
        await self.login()
        moved=await self.client.patch(f'/api/projects/{project["id"]}',json={'expected_revision':latest['revision'],'leader_id':self.outsider_id,'reason':'正式移交'})
        self.assertEqual(moved.status_code,200,moved.text)
        await self.login('leader')
        self.assertEqual((await self.client.get(snapshot_url)).status_code,404)
        self.assertEqual((await self.client.post(f'/api/billing-versions/{version["id"]}/reveal',json={'purpose':'移交后猜测'})).status_code,404)

    async def test_default_profile_and_database_immutable_records(self):
        from sqlalchemy.exc import IntegrityError
        await self.login()
        first,version=await self.verified_profile()
        second,_=await self.verified_profile()
        for profile in (first,second):
            response=await self.client.patch(f'/api/billing-profiles/{profile["id"]}',json={'expected_revision':profile['revision'],'is_default':True})
            self.assertEqual(response.status_code,200,response.text)
        with SessionLocal() as db:
            self.assertEqual(db.query(models.BillingProfile).filter_by(customer_id=self.customer_id,is_default=True,is_active=True).count(),1)
        for sql in (f"UPDATE customer_billing_versions SET public_json='{{}}' WHERE id={version['id']}",f"DELETE FROM customer_billing_versions WHERE id={version['id']}","DELETE FROM audit_events"):
            with engine.begin() as connection:
                with self.assertRaises(IntegrityError):
                    connection.exec_driver_sql(sql)

    async def test_project_requests_execute_once_and_reject_stale_project(self):
        await self.login()
        project=await self.new_project()
        await self.login('leader')
        request=await self.client.post(f'/api/projects/{project["id"]}/change-requests',json={'expected_revision':1,'kind':'transfer','target_leader_id':self.outsider_id,'reason':'项目交接'})
        self.assertEqual(request.status_code,200,request.text)
        await self.login('number')
        url=f'/api/project-change-requests/{request.json()["id"]}/review'
        decision={'expected_revision':1,'decision':'approve','reason':'核对人员资格和交接清单'}
        self.assertEqual((await self.client.post(url,json=decision)).status_code,200)
        self.assertEqual((await self.client.post(url,json=decision)).status_code,409)
        self.assertEqual((await self.client.get(f'/api/projects/{project["id"]}')).json()['leader_id'],self.outsider_id)
        await self.login('leader')
        self.assertEqual((await self.client.get(f'/api/projects/{project["id"]}')).status_code,403)
        await self.login()
        project=await self.new_project()
        await self.login('leader')
        request=await self.client.post(f'/api/projects/{project["id"]}/change-requests',json={'expected_revision':1,'kind':'transfer','target_leader_id':self.outsider_id,'reason':'申请后项目变化'})
        changed=await self.client.patch(f'/api/projects/{project["id"]}',json={'expected_revision':1,'contract_no':'新合同'})
        self.assertEqual(changed.status_code,200,changed.text)
        await self.login()
        self.assertEqual((await self.client.post(f'/api/project-change-requests/{request.json()["id"]}/review',json=decision)).status_code,409)

    async def test_project_customer_correction_preserves_historical_identity(self):
        await self.login()
        project=await self.new_project()
        customer=await self.client.post('/api/customers',json={'name':'纠错目标','tax_id':'CORRECTION-TAX','type':'enterprise'})
        customer_id=customer.json()['id']
        with SessionLocal() as db:
            db.get(models.Customer,customer_id).identity_status='confirmed';db.commit()
        data={'expected_revision':1,'customer_id':customer_id,'identity_checked':True,'reason':'已核对项目合同购买主体'}
        await self.login('finance')
        self.assertEqual((await self.client.post(f'/api/projects/{project["id"]}/correct-customer',json=data)).status_code,403)
        await self.login()
        self.assertEqual((await self.client.post(f'/api/projects/{project["id"]}/correct-customer',json=data)).status_code,200)
        latest=(await self.client.get(f'/api/projects/{project["id"]}')).json()
        self.assertEqual(latest['customer_id'],customer_id)
        self.assertEqual(latest['customer_name'],project['customer_name'])
        self.assertEqual((await self.client.post(f'/api/projects/{project["id"]}/correct-customer',json=data)).status_code,409)

    def test_eight_character_setup_in_fresh_isolated_database(self):
        import subprocess
        with tempfile.TemporaryDirectory(prefix='password-initializer-') as directory:
            env = os.environ.copy()
            env['PYTHONIOENCODING'] = 'utf-8'
            env['FIRM_MANAGER_DATA_DIR'] = directory
            env['PYTHONPATH'] = '/Users/aowu/Documents/同步盘/项目开发/事务所编号管理软件/代码/backend' + os.pathsep + str(Path(__file__).resolve().parents[1])
            result = subprocess.run([sys.executable, '-c', "\nimport asyncio\nimport httpx\nimport main\nfrom database import engine\nasync def check():\n    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=main.app, client=('127.0.0.1', 1234)), base_url='http://testserver') as client:\n        data={'admin_username':'policy-admin','admin_password':'Abcd1234','admin_real_name':'测试','fiscal_year':2026,\n              'firms':[{'name':'隔离所','report_types':[{'report_type':'审计','template':'隔离{yyyy}{nnnn}'}]}]}\n        for invalid in ('Abcd123', 'abcdefgh', '12345678'):\n            response=await client.post('/api/setup', json={**data, 'admin_password':invalid})\n            assert response.status_code==422, response.text\n        response=await client.post('/api/setup',json=data)\n        assert response.status_code==200, response.text\n        response=await client.post('/api/auth/login',json={'username':'policy-admin','password':'Abcd1234'})\n        assert response.status_code==200, response.text\nasyncio.run(check())\nengine.dispose()\n"], env=env, capture_output=True, text=True, encoding='utf-8', timeout=30)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)

    def test_login_password_policy_accepts_eight_and_requires_letters_digits(self):
        from fastapi import HTTPException
        from auth import verify_password
        for value in ('Abcd1234', 'A1' + 'x' * 70, 'Abcd1234!'):
            self.assertTrue(verify_password(value, get_password_hash(value)))
        for value in ('Abcd123', 'abcdefgh', '12345678', '中文密码1234', 'A1' + 'x' * 71, 'Abcd1234' + '中' * 22):
            with self.subTest(value=value), self.assertRaises(HTTPException) as caught:
                get_password_hash(value)
            self.assertEqual(caught.exception.status_code, 422)
            self.assertIn('字母和数字', caught.exception.detail)

    async def test_eight_character_password_create_change_reset_and_legacy_login(self):
        await self.login()
        payload = {'username': 'policy-user', 'real_name': '密码规则测试', 'password': 'Abcd1234',
                   'roles': ['clerk'], 'is_practitioner': False, 'special_grants': []}
        for invalid in ('abcdefgh', '12345678', 'Abcd123'):
            response = await self.client.post('/api/users', json={**payload, 'password': invalid})
            self.assertEqual(response.status_code, 422, response.text)
        created = await self.client.post('/api/users', json=payload)
        self.assertEqual(created.status_code, 200, created.text)
        user_id = created.json()['id']
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=main.app), base_url='http://testserver') as user:
            response = await user.post('/api/auth/login', json={'username': 'policy-user', 'password': 'Abcd1234'})
            self.assertEqual(response.status_code, 200, response.text)
            user.headers['X-CSRF-Token'] = user.cookies['firm_csrf']
            for invalid in ('abcdefgh', '12345678', 'Abcd123'):
                response = await user.post('/api/auth/change-password', json={'current_password': 'Abcd1234', 'new_password': invalid})
                self.assertEqual(response.status_code, 422, response.text)
            response = await user.post('/api/auth/change-password', json={'current_password': 'Abcd1234', 'new_password': 'Change12'})
            self.assertEqual(response.status_code, 200, response.text)
            response = await user.post('/api/auth/login', json={'username': 'policy-user', 'password': 'Change12'})
            self.assertEqual(response.status_code, 200, response.text)
            user.headers['X-CSRF-Token'] = user.cookies['firm_csrf']
            current = (await self.client.get('/api/users')).json()
            revision = next(item['permission_revision'] for item in current if item['id'] == user_id)
            for invalid in ('abcdefgh', '12345678'):
                response = await self.client.put(f'/api/users/{user_id}', json={'expected_revision': revision, 'password': invalid})
                self.assertEqual(response.status_code, 422, response.text)
            response = await self.client.put(f'/api/users/{user_id}', json={'expected_revision': revision, 'password': 'Reset123'})
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual((await user.get('/api/auth/me')).status_code, 401)
            response = await user.post('/api/auth/login', json={'username': 'policy-user', 'password': 'Reset123'})
            self.assertEqual(response.status_code, 200, response.text)
            # Existing passwords are verified as stored, without an upgrade reset.
            from auth import pwd_context
            with SessionLocal() as db:
                db.get(models.User, user_id).hashed_password = pwd_context.hash('legacyonlyletters')
                db.commit()
            response = await user.post('/api/auth/login', json={'username': 'policy-user', 'password': 'legacyonlyletters'})
            self.assertEqual(response.status_code, 200, response.text)

    async def test_password_change_revokes_other_device(self):
        await self.login("leader")
        other = httpx.AsyncClient(transport=httpx.ASGITransport(app=main.app), base_url="http://testserver")
        try:
            await other.post("/api/auth/login", json={"username": "leader", "password": "Test-password-2026"})
            response = await self.client.post("/api/auth/change-password", json={"current_password": "Test-password-2026", "new_password": "New-test-password-2026"})
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual((await other.get("/api/auth/me")).status_code, 401)
        finally:
            await other.aclose()
            hashed = get_password_hash("Test-password-2026")
            with SessionLocal() as db:
                db.query(models.User).filter_by(username="leader").one().hashed_password = hashed
                db.commit()

    async def test_cookie_flags_and_csrf(self):
        response = await self.login()
        cookie = response.headers.get_list("set-cookie")[0].lower()
        self.assertIn("httponly", cookie)
        self.assertIn("samesite=lax", cookie)
        self.assertNotIn("; secure", cookie)
        del self.client.headers["X-CSRF-Token"]
        self.assertEqual((await self.client.put("/api/projects/1", json={})).status_code, 403)

    async def test_cross_origin_login_and_write_denied(self):
        response = await self.client.post("/api/auth/login", headers={"Origin": "http://evil.example"}, json={"username": "admin", "password": "Test-password-2026"})
        self.assertEqual(response.status_code, 403)

    async def test_chunked_body_limit(self):
        async def body():
            for _ in range(33):
                yield b"x" * 65536
        response = await self.client.post("/api/auth/login", content=body())
        self.assertEqual(response.status_code, 413)

    async def test_logout_revokes_cookie(self):
        await self.login()
        original = self.client.cookies["firm_session"]
        self.assertEqual((await self.client.post("/api/auth/logout")).status_code, 200)
        self.client.cookies.set("firm_session", original)
        self.assertEqual((await self.client.get("/api/auth/me")).status_code, 401)

    async def test_admin_reset_revokes_existing_session(self):
        await self.login("outsider")
        with SessionLocal() as db:
            user = db.query(models.User).filter_by(username="outsider").one()
            user.session_version += 1
            db.commit()
        self.assertEqual((await self.client.get("/api/auth/me")).status_code, 401)

    async def test_spoofed_host_and_forwarded_origin(self):
        self.assertEqual((await self.client.get("/api/setup/status", headers={"Host": "evil.example"})).status_code, 400)
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=main.app, client=("203.0.113.5", 12345)), base_url="http://testserver") as remote:
            response = await remote.post("/api/setup", headers={"X-Forwarded-For": "127.0.0.1"}, json={"admin_username": "other", "admin_password": "Test-password-2026", "admin_real_name": "管理", "fiscal_year": 2026, "firms": []})
            self.assertEqual(response.status_code, 403)

    async def test_request_id_and_no_store(self):
        response = await self.client.get("/api/projects")
        self.assertTrue(response.headers["X-Request-ID"])
        self.assertEqual(response.headers["Cache-Control"], "no-store")
        self.assertEqual(response.json()["detail"]["request_id"], response.headers["X-Request-ID"])

    async def export_job(self,**changes):
        import secrets
        payload={'export_type':'projects','filters':{'fiscal_year':2026},'columns':['project_id','customer_name']}
        payload.update(changes)
        response=await self.client.post('/api/export-jobs',json=payload,headers={'Idempotency-Key':secrets.token_hex(16)})
        self.assertEqual(response.status_code,202,response.text)
        return response.json()

    async def test_export_idempotency_limits_scope_and_download_revalidation(self):
        import secrets
        import exports
        from io import BytesIO
        from openpyxl import load_workbook
        await self.login('number')
        denied=await self.client.post('/api/export-jobs',json={'filters':{'fiscal_year':2026},'columns':['contract_amount']},headers={'Idempotency-Key':secrets.token_hex(16)})
        self.assertEqual(denied.status_code,403)
        key=secrets.token_hex(16);payload={'filters':{'fiscal_year':2026},'columns':['project_id','customer_name']}
        response=await self.client.post('/api/export-jobs',json=payload,headers={'Idempotency-Key':key});self.assertEqual(response.status_code,202,response.text)
        first=response.json()
        retry=await self.client.post('/api/export-jobs',json=payload,headers={'Idempotency-Key':key})
        self.assertEqual(first['id'],retry.json()['id'])
        changed=await self.client.post('/api/export-jobs',json={**payload,'columns':['project_id']},headers={'Idempotency-Key':key})
        self.assertEqual(changed.status_code,409)
        second=await self.export_job()
        third=await self.client.post('/api/export-jobs',json=payload,headers={'Idempotency-Key':secrets.token_hex(16)})
        self.assertEqual(third.status_code,429)
        self.assertEqual((await self.client.post('/api/export-jobs/'+second['id']+'/cancel')).status_code,200)
        claimed=await asyncio.to_thread(exports.claim);self.assertEqual(claimed,first['id'])
        await asyncio.to_thread(exports.generate,claimed)
        state=(await self.client.get('/api/export-jobs/'+claimed)).json();self.assertEqual(state['status'],'succeeded',state)
        downloaded=await self.client.get('/api/export-jobs/'+claimed+'/download');self.assertEqual(downloaded.status_code,200,downloaded.text[:100])
        workbook=load_workbook(BytesIO(downloaded.content),read_only=True)
        self.assertEqual(list(workbook['资料'].values)[0],('项目ID','客户名称'));workbook.close()
        await self.login('outsider');self.assertEqual((await self.client.get('/api/export-jobs/'+claimed+'/download')).status_code,404)
        await self.login('number')
        with SessionLocal() as db:
            project=db.get(models.Project,1);project.leader_id=self.outsider_id;db.commit()
        self.assertEqual((await self.client.get('/api/export-jobs/'+claimed+'/download')).status_code,403)
        with SessionLocal() as db:db.get(models.Project,1).leader_id=self.leader_id;db.commit()

    async def test_export_worker_process_cancel_restart_ttl_and_formula_safety(self):
        import exports,subprocess,time
        from unittest.mock import patch
        from io import BytesIO
        from openpyxl import load_workbook
        await self.login()
        project=await self.new_project(customer_name='=危险公式')
        with SessionLocal() as db:
            db.get(models.Project,project['id']).customer_name='=危险公式';db.commit()
        job=await self.export_job(filters={'fiscal_year':2026,'search':project['project_id']})
        worker=subprocess.Popen([sys.executable,str(Path(main.__file__).with_name('export_worker.py'))],stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
        try:
            for _ in range(100):
                state=(await self.client.get('/api/export-jobs/'+job['id'])).json()
                if state['status'] not in ('queued','running'):break
                await asyncio.sleep(.1)
            self.assertEqual(state['status'],'succeeded',state)
            downloaded=await self.client.get('/api/export-jobs/'+job['id']+'/download')
            workbook=load_workbook(BytesIO(downloaded.content),read_only=True)
            self.assertEqual(list(workbook['资料'].values)[1][1],"'=危险公式");workbook.close()
            second=subprocess.run([sys.executable,str(Path(main.__file__).with_name('export_worker.py'))],timeout=5)
            self.assertEqual(second.returncode,0)
        finally:worker.terminate();worker.wait(timeout=5)
        with SessionLocal() as db:db.get(models.ExportJob,job['id']).expires_at=datetime(2000,1,1);db.commit()
        exports.cleanup();self.assertEqual((await self.client.get('/api/export-jobs/'+job['id']+'/download')).status_code,410)
        queued=await self.export_job();running=exports.claim();self.assertEqual(running,queued['id'])
        self.assertEqual((await self.client.post('/api/export-jobs/'+running+'/cancel')).status_code,200)
        await asyncio.to_thread(exports.generate,running)
        self.assertEqual((await self.client.get('/api/export-jobs/'+running)).json()['status'],'cancelled')
        interrupted=await self.export_job();exports.claim();exports.cleanup(restart=True)
        state=(await self.client.get('/api/export-jobs/'+interrupted['id'])).json()
        self.assertEqual(state['error_code'],'interrupted')
        budget=await self.export_job();exports.claim()
        with patch.object(exports,'MAX_ROWS',0):await asyncio.to_thread(exports.generate,budget['id'])
        self.assertEqual((await self.client.get('/api/export-jobs/'+budget['id'])).json()['error_code'],'row_limit_exceeded')

    async def test_encrypted_backup_restore_preserves_identity_and_requires_reconciliation(self):
        from contextlib import closing
        from backups import FILES,validate_restored
        from backup_crypto import restore
        from schema_init import SCHEMA_VERSION
        await self.login()
        result=await self.client.post('/api/system/backup',json={'password':'Backup-test-password-2026'})
        self.assertEqual(result.status_code,200,result.text[:100])
        self.assertNotIn(b'Test-password-2026',result.content)
        with tempfile.TemporaryDirectory() as temporary:
            target=Path(temporary)/'restored'
            restore(result.content,'Backup-test-password-2026',target,'YMH-FMC',SCHEMA_VERSION,'db.sqlite',FILES,validate_restored)
            self.assertEqual((target/'billing.key').read_bytes(),(Path(TEST_DIR.name)/'billing.key').read_bytes())
            self.assertEqual((target/'device-identity.json').read_bytes(),(Path(TEST_DIR.name)/'device-identity.json').read_bytes())
            with closing(sqlite3.connect(target/'db.sqlite')) as db:
                self.assertEqual(db.execute("SELECT value FROM app_settings WHERE key='restore_hold'").fetchone()[0],'{"required":true}')
                self.assertEqual(db.execute('SELECT count(*) FROM auth_sessions WHERE revoked_at IS NULL').fetchone()[0],0)
                self.assertEqual(db.execute('PRAGMA integrity_check').fetchone()[0],'ok')
            for blob,password in ((result.content,'wrong-password-2026'),(result.content[:-1]+bytes([result.content[-1]^1]),'Backup-test-password-2026')):
                with self.assertRaises(ValueError):restore(blob,password,Path(temporary)/'bad','YMH-FMC',SCHEMA_VERSION,'db.sqlite',FILES,validate_restored)
            with self.assertRaises(ValueError):restore(result.content,'Backup-test-password-2026',target,'YMH-FMC',SCHEMA_VERSION,'db.sqlite',FILES,validate_restored)

    async def test_export_batches_use_one_snapshot_and_sensitive_requires_separate_grant(self):
        import exports,secrets
        from unittest.mock import patch
        from io import BytesIO
        from openpyxl import load_workbook
        await self.login()
        prefix=secrets.token_hex(6)
        with SessionLocal() as db:
            rows=[models.Project(project_id=prefix+'-'+str(i),firm=prefix,report_type='审计',report_year=2026,
                fiscal_year=2026,customer_name='原名称',customer_id=self.customer_id,leader_id=self.leader_id) for i in range(550)]
            db.add_all(rows);db.flush();last_id=rows[-1].id;db.commit()
        job=await self.export_job(filters={'fiscal_year':2026,'firm':prefix})
        self.assertEqual(exports.claim(),job['id'])
        original=exports.check;calls=0
        def concurrent_change(*args):
            nonlocal calls
            calls+=1
            if calls==3:
                with SessionLocal() as db:db.get(models.Project,last_id).customer_name='并发新名称';db.commit()
            original(*args)
        with patch.object(exports,'check',side_effect=concurrent_change):await asyncio.to_thread(exports.generate,job['id'])
        downloaded=await self.client.get('/api/export-jobs/'+job['id']+'/download')
        self.assertEqual(downloaded.status_code,200,downloaded.text[:100])
        book=load_workbook(BytesIO(downloaded.content),read_only=True)
        values=list(book['资料'].values);self.assertEqual(len(values),551)
        self.assertEqual({row[1] for row in values[1:]},{'原名称'});book.close()
        await self.login('finance')
        response=await self.client.post('/api/export-jobs',json={'export_type':'billing_sensitive','purpose':'测试','confirm_sensitive':True},headers={'Idempotency-Key':secrets.token_hex(16)})
        self.assertEqual(response.status_code,403)
        await self.login()
        response=await self.client.post('/api/export-jobs',json={'export_type':'billing_sensitive','purpose':'测试'},headers={'Idempotency-Key':secrets.token_hex(16)})
        self.assertEqual(response.status_code,422)
        sensitive=await self.export_job(export_type='billing_sensitive',filters=None,columns=[],purpose='完整资料验收',confirm_sensitive=True)
        self.assertEqual(exports.claim(),sensitive['id']);await asyncio.to_thread(exports.generate,sensitive['id'])
        self.assertEqual((await self.client.get('/api/export-jobs/'+sensitive['id'])).json()['status'],'succeeded')

    async def test_restore_hold_blocks_numbering_until_complete_external_reconciliation(self):
        import json
        await self.login()
        with SessionLocal() as db:
            db.merge(models.AppSetting(key='restore_hold',value='{"required":true}'));db.commit()
        try:
            denied=await self.client.put('/api/projects/1',json={'expected_revision':1,'customer_name':'禁止写入'})
            self.assertEqual(denied.status_code,409)
            state=(await self.client.get('/api/system/backup/reconciliation')).json();self.assertTrue(state['required'])
            payload={'rules':[{'id':r['id'],'actual_last_sequence':r['current_sequence']+100} for r in state['rules']], 'confirmed_external_records':False,'reason':'已向外部档案核对已交付最大序号'}
            self.assertEqual((await self.client.post('/api/system/backup/reconcile',json=payload)).status_code,422)
            payload['confirmed_external_records']=True
            self.assertEqual((await self.client.post('/api/system/backup/reconcile',json=payload)).status_code,200)
            self.assertFalse((await self.client.get('/api/system/backup/reconciliation')).json()['required'])
        finally:
            with SessionLocal() as db:db.query(models.AppSetting).filter_by(key='restore_hold').delete();db.commit()

    async def test_export_readonly_download_slot_and_disk_failure(self):
        import exports
        from unittest.mock import patch
        from collections import namedtuple
        await self.login()
        with patch.object(main,'license_required',return_value=True),patch.object(main,'license_status',return_value={'allowed':False,'reason':'只读'}):
            job=await self.export_job();self.assertEqual(exports.claim(),job['id']);await asyncio.to_thread(exports.generate,job['id'])
            exports.download_slot.acquire()
            try:
                response=await self.client.get('/api/export-jobs/'+job['id']+'/download')
                self.assertEqual(response.status_code,429);self.assertEqual(response.headers['Retry-After'],'5')
                from unittest.mock import patch
                with patch('backups.create') as compression:
                    busy=await self.client.post('/api/system/backup',json={'password':'Test-backup-password'})
                    self.assertEqual(busy.status_code,429);compression.assert_not_called()
                self.assertEqual((await self.client.post('/api/diagnostics/export')).status_code,429)
            finally:exports.download_slot.release()
            self.assertEqual((await self.client.get('/api/export-jobs/'+job['id']+'/download')).status_code,200)
            # A completed response releases its slot for the next authenticated retry.
            self.assertEqual((await self.client.get('/api/export-jobs/'+job['id']+'/download')).status_code,200)
        job=await self.export_job();exports.claim()
        usage=namedtuple('Usage','total used free')
        with patch.object(exports.shutil,'disk_usage',return_value=usage(100,100,0)):await asyncio.to_thread(exports.generate,job['id'])
        state=(await self.client.get('/api/export-jobs/'+job['id'])).json();self.assertEqual(state['error_code'],'disk_space_low')
        self.assertFalse((exports.DIRECTORY/(job['id']+'.partial.xlsx')).exists())

    async def test_ten_export_users_fifo_single_generation_and_global_limit(self):
        import exports,secrets
        from unittest.mock import patch
        await self.login()
        prefix='queue-'+secrets.token_hex(4)
        with SessionLocal() as db:
            hashed=get_password_hash('Test-password-2026')
            users=[models.User(username=prefix+str(i),real_name='队列用户',fiscal_year=2026,hashed_password=hashed,
                role_records=[models.UserRoleRecord(role_code='number_manager')]) for i in range(10)]
            db.add_all(users);db.commit()
        async def submit(i):
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=main.app),base_url='http://testserver') as client:
                await client.post('/api/auth/login',json={'username':prefix+str(i),'password':'Test-password-2026'})
                response=await client.post('/api/export-jobs',json={'filters':{'fiscal_year':2026},'columns':['project_id']},
                    headers={'X-CSRF-Token':client.cookies['firm_csrf'],'Idempotency-Key':secrets.token_hex(16)})
                self.assertEqual(response.status_code,202,response.text);return response.json()['id']
        jobs=await asyncio.gather(*(submit(i) for i in range(10)))
        with SessionLocal() as db:
            expected=db.query(models.ExportJob).filter(models.ExportJob.id.in_(jobs)).order_by(models.ExportJob.created_at,models.ExportJob.id).first().id
        self.assertEqual(exports.claim(),expected);self.assertIsNone(exports.claim())
        with SessionLocal() as db:
            self.assertEqual(db.query(models.ExportJob).filter(models.ExportJob.id.in_(jobs),models.ExportJob.status=='queued').count(),9)
        with patch.object(exports,'MAX_PENDING',10):
            response=await self.client.post('/api/export-jobs',json={'filters':{'fiscal_year':2026},'columns':['project_id']},headers={'Idempotency-Key':secrets.token_hex(16)})
            self.assertEqual(response.status_code,429)
        for job in jobs:self.assertEqual((await self.client.post('/api/export-jobs/'+job+'/cancel')).status_code,200)
        await asyncio.to_thread(exports.generate,expected)

    async def test_role_matrix_project_scope_money_edit_and_numbering(self):
        import secrets
        for account,readable,money,editable,issuable in [
            ('admin',True,True,True,True),('number',True,False,True,True),
            ('finance',True,True,False,False),('leader',True,True,True,True),
            ('outsider',True,False,False,False),('combined',True,False,True,True),
            ('clerk',False,False,False,False)]:
            with self.subTest(account=account):
                await self.login();project=await self.new_project(member_ids=[self.outsider_id],contract_amount=99)
                await self.login(account);url='/api/projects/'+str(project['id'])
                detail=await self.client.get(url);self.assertEqual(detail.status_code,200 if readable else 403)
                listing=(await self.client.get('/api/projects',params={'fiscal_year':2026,'search':project['project_id']})).json()
                self.assertEqual(any(p['id']==project['id'] for p in listing['items']),readable)
                if readable:self.assertEqual(detail.json()['contract_amount'],99 if money else None)
                edit=await self.client.patch(url,json={'expected_revision':1,'priority':'高'})
                self.assertEqual(edit.status_code,200 if editable else 403,edit.text)
                revision=2 if editable else 1
                issue=await self.client.post(url+'/generate-report-no',json={'expected_revision':revision},headers={'Idempotency-Key':secrets.token_hex(16)})
                self.assertEqual(issue.status_code,200 if issuable else 403,issue.text)

    def test_foreign_keys_and_full_durability(self):
        with engine.connect() as connection:
            self.assertEqual(connection.exec_driver_sql("PRAGMA foreign_keys").scalar(), 1)
            self.assertEqual(connection.exec_driver_sql("PRAGMA synchronous").scalar(), 2)

    async def test_mounted_static_gzip_hash_cache_and_html_revalidation(self):
        from fastapi import FastAPI
        from fastapi.staticfiles import StaticFiles
        from fastapi.responses import HTMLResponse
        from static_delivery import StaticDelivery
        with tempfile.TemporaryDirectory() as temporary:
            Path(temporary,'index-AbCd1234.js').write_text('const value="'+'x'*1000+'";')
            application=FastAPI();application.add_middleware(StaticDelivery)
            application.mount('/assets',StaticFiles(directory=temporary))
            @application.get('/')
            def index():return HTMLResponse('<html>'+('x'*1000)+'</html>')
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=application),base_url='http://testserver') as client:
                asset=await client.get('/assets/index-AbCd1234.js')
                self.assertEqual(asset.headers['Content-Encoding'],'gzip')
                self.assertEqual(asset.headers['Cache-Control'],'public, max-age=31536000, immutable')
                self.assertEqual((await client.get('/')).headers['Cache-Control'],'no-cache')

    def test_unknown_schema_is_preserved(self):
        path = Path(TEST_DIR.name) / "unknown.sqlite"
        with sqlite3.connect(path) as db:
            db.execute("CREATE TABLE old_data(value TEXT)")
            db.execute("INSERT INTO old_data VALUES('keep')")
        old = create_engine(f"sqlite:///{path}")
        with self.assertRaises(RuntimeError):
            initialize_schema(old)
        with sqlite3.connect(path) as db:
            self.assertEqual(db.execute("SELECT value FROM old_data").fetchone()[0], "keep")
        old.dispose()


    async def test_diagnostics_admin_can_trace_practitioner_denial(self):
        import io, json, zipfile
        await self.login("leader")
        denied = await self.client.get('/api/users')
        self.assertEqual(denied.status_code,403)
        request_id = denied.json()['detail']['request_id']
        self.assertEqual(denied.json()['detail']['category'],'access_restriction')
        for username in ('leader','outsider','number','finance','clerk','combined'):
            await self.login(username)
            self.assertEqual((await self.client.get('/api/diagnostics/logs')).status_code,403)
            self.assertEqual((await self.client.post('/api/diagnostics/export')).status_code,403)
        await self.login('admin')
        filters={'request_id':request_id,'actor_id':self.leader_id}
        response=await self.client.get('/api/diagnostics/logs',params=filters)
        rows=response.json()['items']
        self.assertEqual(len(rows),1,response.text)
        self.assertEqual(rows[0]['status_code'],403)
        self.assertEqual(rows[0]['actor_id'],self.leader_id)
        self.assertIn('timestamp',rows[0])
        exported=await self.client.post('/api/diagnostics/export',params=filters)
        self.assertEqual(exported.status_code,200,exported.text if exported.status_code!=200 else '')
        with zipfile.ZipFile(io.BytesIO(exported.content)) as archive:
            records=[json.loads(line) for line in archive.read('logs.jsonl').splitlines()]
            self.assertEqual(records,rows)
            self.assertFalse(json.loads(archive.read('summary.json'))['truncated'])
        self.assertNotIn('Test-password',str(rows))
        self.assertNotIn('firm_session',str(rows))
        self.assertEqual((await self.client.post('/api/diagnostics/export',headers={'X-CSRF-Token':'bad'})).status_code,403)
        self.assertEqual((await self.client.get('/api/diagnostics/logs',params={'since':'2026-01-01','until':'2026-10-01'})).status_code,422)

    async def test_log_disk_refusal_does_not_reverse_committed_project(self):
        from unittest.mock import patch
        import runtime_log
        await self.login()
        before=runtime_log.health.snapshot()['log_drops']
        with patch.object(runtime_log.logger,'info',side_effect=PermissionError('test disk refusal')):
            project=await self.new_project()
        with SessionLocal() as db:self.assertIsNotNone(db.get(models.Project,project['id']))
        state=runtime_log.health.snapshot()
        self.assertGreater(state['log_drops'],before)
        self.assertEqual(state['last_log_failure_code'],'log_write_failed')
        runtime_log.event('probe.recovered')
        self.assertFalse(runtime_log.health.pending)

    async def test_chinese_unexpected_exception_and_safe_diagnostics(self):
        from fastapi import FastAPI, Request
        from http_security import install_http_boundary
        from diagnostics import selected, Filters
        from diagnostic_reader import read_records, window
        from database import DATA_DIR
        probe=FastAPI()
        install_http_boundary(probe,'FIRM_MANAGER',80)
        @probe.get('/api/probe')
        def fail(request:Request):
            request.state.actor_id=self.leader_id
            raise RuntimeError('SECRET-SQL bank=622222 password=should-never-leak')
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=probe),base_url='http://testserver') as client:
            response=await client.get('/api/probe')
            self.assertEqual(response.status_code,500,response.text)
            detail=response.json()['detail']
            self.assertEqual(detail['category'],'system_error')
            self.assertIn('服务处理失败',detail['message'])
            request_id=detail['request_id']
            self.assertEqual(response.headers['X-Request-ID'],request_id)
            missing=await client.get('/api/missing')
            self.assertIn('不存在',missing.json()['detail']['message'])
        since,until=window(None,None)
        rows,_=read_records(DATA_DIR/'logs',('operations.jsonl*',),since,until,{'request_id':request_id})
        self.assertEqual(len(rows),1)
        self.assertEqual(rows[0]['exception_type'],'RuntimeError')
        self.assertIn('fail',rows[0]['exception_location'])
        self.assertNotIn('SECRET-SQL',str(rows)+response.text)
        self.assertNotIn('622222',str(rows)+response.text)

    def test_diagnostic_reader_timezone_and_projection(self):
        import json
        from datetime import timezone
        from diagnostic_reader import window, read_records
        since,until=window(datetime.fromisoformat('2026-10-02T08:00:00+08:00'),datetime.fromisoformat('2026-10-02T09:00:00+08:00'))
        self.assertEqual(since.hour,0)
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'operations.jsonl'
            path.write_text(json.dumps({'timestamp':'2026-10-02T00:30:00Z','event':'probe','password':'SECRET','bank_account':'622222','headers':{'Cookie':'secret'}})+'\n{bad json}\n')
            rows,_=read_records(Path(directory),('operations.jsonl*',),since,until,{})
            self.assertEqual(rows,[{'timestamp':'2026-10-02T00:30:00Z','event':'probe'}])

    async def task_fixture(self, customer_id=None):
        await self.login()
        customer_id = customer_id or self.customer_id
        with SessionLocal() as db: customer_name=db.get(models.Customer,customer_id).name
        project = await self.new_project(contract_amount=1000,customer_id=customer_id,customer_name=customer_name)
        profile, version = await self.verified_profile(customer_id=customer_id)
        with SessionLocal() as db:
            assignee = db.query(models.User).filter_by(username='clerk').one().id
        data = {'project_id': project['id'], 'customer_id': customer_id,
                'billing_profile_id': profile['id'], 'billing_version_id': version['id'],
                'assignee_id': assignee, 'amount': '123.45', 'note': '按本次任务办理'}
        created = await self.client.post('/api/billing-tasks', json=data,
            headers={'Idempotency-Key': 'task-create-'+str(project['id'])})
        self.assertEqual(created.status_code, 200, created.text)
        return created.json(), data, project, profile, version

    async def test_billing_task_assignment_minimal_scope_and_finance_confirmation(self):
        task, data, project, profile, version = await self.task_fixture()
        await self.login('outsider')
        self.assertEqual((await self.client.get('/api/billing-tasks/'+task['id'])).status_code, 404)
        await self.login('clerk')
        rows = (await self.client.get('/api/billing-tasks', params={'owner': self.admin_id})).json()
        self.assertIn(task['id'], [item['id'] for item in rows['items']])
        for field in ['contract_amount', 'received_amount', 'bank_account', 'contact_name']:
            self.assertNotIn(field, str(rows))
        reveal = await self.client.post('/api/billing-tasks/'+task['id']+'/reveal',json={'purpose':'核对本次开票'})
        self.assertEqual(reveal.status_code, 200, reveal.text)
        self.assertEqual(reveal.json()['fields']['bank_account'],'0012345678901234')
        self.assertEqual(reveal.headers['Cache-Control'],'no-store')
        self.assertEqual((await self.client.post('/api/billing-versions/'+str(version['id'])+'/reveal',json={'purpose':'猜旧接口'})).status_code,404)
        for path in ['/api/diagnostics/logs','/api/diagnostics/summary']:
            self.assertEqual((await self.client.get(path)).status_code,403)
        result = {'expected_revision': task['revision'], 'amount':'123.45','invoice_number':'TASK-INV-'+task['id'][:8], 'invoice_date':'2026-10-02'}
        url = '/api/billing-tasks/'+task['id']+'/submit-result'
        submitted = await self.client.post(url,json=result,headers={'Idempotency-Key':'submit-'+task['id']})
        self.assertEqual(submitted.status_code,200,submitted.text)
        again = await self.client.post(url,json=result,headers={'Idempotency-Key':'submit-'+task['id']})
        self.assertEqual(again.status_code,200,again.text)
        with SessionLocal() as db:
            self.assertEqual(db.query(models.FinancialEntry).filter_by(project_id=project['id']).count(),0)
        await self.login('finance')
        confirm={'expected_revision':submitted.json()['revision'],'expected_project_revision':project['revision'],'reason':'财务核对实际凭证'}
        url='/api/billing-tasks/'+task['id']+'/confirm'
        replies=await asyncio.gather(*(self.client.post(url,json=confirm,headers={'Idempotency-Key':'confirm-'+task['id']}) for _ in range(3)))
        self.assertTrue(all(r.status_code==200 for r in replies),[r.text for r in replies])
        self.assertEqual(len({r.json()['financial_entry_id'] for r in replies}),1)
        with SessionLocal() as db:
            self.assertEqual(db.query(models.FinancialEntry).filter_by(project_id=project['id']).count(),1)
        await self.login('clerk')
        self.assertEqual((await self.client.post('/api/billing-tasks/'+task['id']+'/reveal',json={'purpose':'完成后猜取'})).status_code,409)

    async def test_billing_task_cross_objects_duplicate_create_and_new_role_permissions(self):
        task,data,project,profile,version=await self.task_fixture()
        repeated=await self.client.post('/api/billing-tasks',json=data,headers={'Idempotency-Key':'task-create-'+str(project['id'])})
        self.assertEqual(repeated.json()['id'],task['id'])
        audits=await self.client.get('/api/audit-events',params={'target_type':'billing_tasks','target_id':task['id']})
        self.assertEqual(audits.status_code,200,audits.text)
        self.assertEqual(audits.json()['total'],1)
        self.assertTrue(audits.json()['items'][0]['request_id'])
        conflict=await self.client.post('/api/billing-tasks',json={**data,'amount':'1.00'},headers={'Idempotency-Key':'task-create-'+str(project['id'])})
        self.assertEqual(conflict.status_code,409)
        other_profile,_=await self.verified_profile()
        self.assertEqual((await self.client.post('/api/billing-tasks',json={**data,'billing_profile_id':other_profile['id']},headers={'Idempotency-Key':'cross-task-'+task['id']})).status_code,422)
        self.assertEqual((await self.client.post('/api/billing-tasks',json={**data,'assignee_id':self.leader_id},headers={'Idempotency-Key':'bad-assignee-'+task['id']})).status_code,422)
        for username in ['number','leader','finance','clerk']:
            await self.login(username)
            self.assertEqual((await self.client.get('/api/diagnostics/logs')).status_code,403)
            self.assertEqual((await self.client.post('/api/system/backup',json={'password':'Test-backup-password'})).status_code,403)
        await self.login('clerk')
        self.assertEqual((await self.client.post('/api/billing-tasks',json=data,headers={'Idempotency-Key':'clerk-create-'+task['id']})).status_code,403)
        self.assertEqual((await self.client.post('/api/billing-tasks/'+task['id']+'/reveal',json={'purpose':'核对','billing_version_id':other_profile['id']})).status_code,422)
        self.assertEqual((await self.client.get('/api/billing-tasks/'+task['id']+'bad')).status_code,404)

    async def test_billing_task_reassign_revoke_and_stale_window(self):
        task,_,_,_,_=await self.task_fixture()
        with SessionLocal() as db:
            other=models.User(username='task-clerk-'+task['id'][:8],real_name='另一后勤',fiscal_year=2026,
                hashed_password=get_password_hash('Test-password-2026'),role_records=[models.UserRoleRecord(role_code='clerk')])
            db.add(other);db.commit(); other_id=other.id;other_name=other.username
        url='/api/billing-tasks/'+task['id']
        await self.login(other_name)
        self.assertEqual((await self.client.get(url)).status_code,404)
        self.assertNotIn(task['id'],[x['id'] for x in (await self.client.get('/api/billing-tasks')).json()['items']])
        self.assertEqual((await self.client.post(url+'/reveal',json={'purpose':'他人任务猜取'})).status_code,404)
        await self.login()
        moved=await self.client.post(url+'/reassign',json={'expected_revision':1,'assignee_id':other_id,'reason':'交接办理'})
        self.assertEqual(moved.status_code,200,moved.text)
        self.assertEqual((await self.client.post(url+'/revoke',json={'expected_revision':1,'reason':'旧窗口操作'})).status_code,409)
        await self.login('clerk')
        self.assertEqual((await self.client.get(url)).status_code,404)
        self.assertEqual((await self.client.post(url+'/reveal',json={'purpose':'旧窗口'})).status_code,404)
        self.assertNotIn(task['id'],[x['id'] for x in (await self.client.get('/api/billing-tasks')).json()['items']])
        await self.login(other_name)
        self.assertEqual((await self.client.post(url+'/reveal',json={'purpose':'本人办理'})).status_code,200)
        await self.login()
        revoked=await self.client.post(url+'/revoke',json={'expected_revision':moved.json()['revision'],'reason':'本次取消'})
        self.assertEqual(revoked.status_code,200)
        await self.login(other_name)
        self.assertEqual((await self.client.post(url+'/reveal',json={'purpose':'撤回后猜取'})).status_code,409)

    async def test_billing_task_version_change_requires_revalidation_and_preserves_history(self):
        task,_,project,profile,version=await self.task_fixture()
        url='/api/billing-tasks/'+task['id']
        await self.login('clerk')
        submission={'expected_revision':1,'amount':'123.45','invoice_number':'HIST-'+task['id'][:8],'invoice_date':'2026-10-02'}
        submitted=await self.client.post(url+'/submit-result',json=submission,headers={'Idempotency-Key':'history-submit-'+task['id']})
        self.assertEqual(submitted.status_code,200)
        self.assertEqual((await self.client.post(url+'/reveal',json={'purpose':'已提交读取'})).status_code,409)
        await self.login()
        changed=await self.client.post(f'/api/billing-profiles/{profile["id"]}/versions',json={'expected_profile_revision':profile['revision'],'fields':{'bank_account':'0099999999999999'}})
        newer=changed.json()
        verified=await self.client.post(f'/api/billing-versions/{newer["id"]}/verify',json={'expected_revision':newer['revision'],'expected_profile_revision':profile['revision']+1,'decision':'approve','verification_note':'新银行资料'})
        self.assertEqual(verified.status_code,200,verified.text)
        confirmation={'expected_revision':2,'expected_project_revision':project['revision'],'reason':'已实际开票，核对旧开户资料'}
        self.assertEqual((await self.client.post(url+'/confirm',json=confirmation,headers={'Idempotency-Key':'historic-confirm-'+task['id']})).status_code,409)
        confirmed=await self.client.post(url+'/confirm',json={**confirmation,'confirm_historical':True},headers={'Idempotency-Key':'historic-confirm-'+task['id']})
        self.assertEqual(confirmed.status_code,200,confirmed.text)
        snapshot=await self.client.get(f'/api/projects/{project["id"]}/finance/invoices/{confirmed.json()["financial_entry_id"]}/billing-snapshot')
        self.assertEqual(snapshot.json()['fields']['bank_account'],'0012345678901234')
        # 普通财务录入仍不可绕过当前版本校验。
        normal={'expected_project_revision':project['revision']+1,'amount':'1.00','occurred_on':'2026-10-02',
            'billing_profile_id':profile['id'],'billing_version_id':version['id'],'expected_profile_revision':verified.json()['profile_revision']}
        self.assertEqual((await self.client.post(f'/api/projects/{project["id"]}/finance/invoices',json=normal,headers={'Idempotency-Key':'old-normal-'+task['id']})).status_code,409)

    async def test_billing_task_needs_review_and_return_keeps_result_history(self):
        task,_,project,profile,version=await self.task_fixture()
        url='/api/billing-tasks/'+task['id']
        with SessionLocal() as db:
            db.get(models.Project,project['id']).is_deleted=True;db.commit()
        await self.login('clerk')
        self.assertEqual((await self.client.get(url)).json()['status'],'needs_review')
        self.assertEqual((await self.client.post(url+'/reveal',json={'purpose':'删除后操作'})).status_code,409)
        await self.login()
        with SessionLocal() as db:
            db.get(models.Project,project['id']).is_deleted=False;db.commit()
        current=(await self.client.get(url)).json()
        revalidated=await self.client.post(url+'/revalidate',json={'expected_revision':current['revision'],'billing_profile_id':profile['id'],'billing_version_id':version['id'],'reason':'核对项目和原资料恢复'})
        self.assertEqual(revalidated.status_code,200,revalidated.text)
        await self.login('clerk')
        result=await self.client.post(url+'/submit-result',json={'expected_revision':revalidated.json()['revision'],'amount':'123.45','invoice_number':'RETURN-'+task['id'][:8],'invoice_date':'2026-10-02'},headers={'Idempotency-Key':'return-submit-'+task['id']})
        self.assertEqual(result.status_code,200,result.text)
        await self.login('finance')
        returned=await self.client.post(url+'/return',json={'expected_revision':result.json()['revision'],'reason':'凭证核对退回'})
        self.assertEqual(returned.status_code,200,returned.text)
        self.assertEqual(len(returned.json()['result_history']),1)
        self.assertIsNone(returned.json()['result'])
        self.assertEqual(returned.json()['status'],'assigned')
        with SessionLocal() as db:
            with self.assertRaises(Exception):
                db.execute(__import__('sqlalchemy').text('DELETE FROM billing_task_results WHERE task_id=:id'),{'id':task['id']});db.commit()
            db.rollback()

    async def test_billing_task_disabled_assignee_and_customer_change(self):
        task,_,_,_,_=await self.task_fixture()
        with SessionLocal() as db:
            clerk=db.query(models.User).filter_by(username='clerk').one();clerk_id=clerk.id
            clerk.is_active=False;db.commit()
        try:
            await self.login()
            self.assertEqual((await self.client.post('/api/billing-tasks/'+task['id']+'/reveal',json={'purpose':'禁用办理人'})).status_code,422)
        finally:
            with SessionLocal() as db:db.get(models.User,clerk_id).is_active=True;db.commit()
        with SessionLocal() as db:
            customer=db.get(models.Customer,self.customer_id);previous=customer.is_active;customer.is_active=False;db.commit()
        try:
            await self.login('clerk')
            self.assertEqual((await self.client.post('/api/billing-tasks/'+task['id']+'/reveal',json={'purpose':'停用客户'})).status_code,409)
            self.assertEqual((await self.client.get('/api/billing-tasks/'+task['id'])).json()['status'],'needs_review')
        finally:
            with SessionLocal() as db:db.get(models.Customer,self.customer_id).is_active=previous;db.commit()

    async def test_billing_task_link_existing_invoice_matches_all_fields(self):
        task,_,project,profile,version=await self.task_fixture()
        reference='LINK-'+task['id'][:8]
        invoice=await self.client.post(f'/api/projects/{project["id"]}/finance/invoices',json={
            'amount':'123.45','occurred_on':'2026-10-02','reference':reference,
            'expected_project_revision':project['revision'],'billing_profile_id':profile['id'],
            'billing_version_id':version['id'],'expected_profile_revision':profile['revision']},
            headers={'Idempotency-Key':'link-manual-'+task['id']})
        self.assertEqual(invoice.status_code,200,invoice.text)
        await self.login('clerk')
        submitted=await self.client.post('/api/billing-tasks/'+task['id']+'/submit-result',json={
            'expected_revision':1,'amount':'123.45','invoice_number':reference,'invoice_date':'2026-10-02'},
            headers={'Idempotency-Key':'link-submit-'+task['id']})
        self.assertEqual(submitted.status_code,200,submitted.text)
        await self.login('finance')
        linked=await self.client.post('/api/billing-tasks/'+task['id']+'/confirm',json={
            'expected_revision':submitted.json()['revision'],'expected_project_revision':project['revision']+1,
            'financial_entry_id':invoice.json()['id'],'reason':'核对已存在手工流水'},headers={'Idempotency-Key':'link-confirm-'+task['id']})
        self.assertEqual(linked.status_code,200,linked.text)
        self.assertEqual(linked.json()['financial_entry_id'],invoice.json()['id'])
        with SessionLocal() as db:
            self.assertEqual(db.query(models.FinancialEntry).filter_by(project_id=project['id']).count(),1)
            self.assertEqual(db.query(models.InvoiceSnapshot).filter_by(financial_entry_id=invoice.json()['id']).count(),1)

    async def test_billing_task_pending_version_and_customer_merge_need_review(self):
        task,_,_,profile,_=await self.task_fixture()
        changed=await self.client.post(f'/api/billing-profiles/{profile["id"]}/versions',json={
            'expected_profile_revision':profile['revision'],'fields':{'bank_account':'0022222222222222'}})
        new=changed.json()
        verified=await self.client.post(f'/api/billing-versions/{new["id"]}/verify',json={
            'expected_revision':new['revision'],'expected_profile_revision':profile['revision']+1,
            'decision':'approve','verification_note':'待办理期间换银行'})
        self.assertEqual(verified.status_code,200)
        await self.login('clerk')
        url='/api/billing-tasks/'+task['id']
        self.assertEqual((await self.client.get(url)).json()['status'],'needs_review')
        self.assertEqual((await self.client.post(url+'/reveal',json={'purpose':'读取旧版本'})).status_code,409)
        await self.login()
        with SessionLocal() as db:
            source=models.Customer(name='任务合并来源',identity_status='confirmed',type='enterprise',tax_id=('TASK-MERGE-'+task['id'][:8]).upper());db.add(source);db.commit();source_id=source.id
        merged_task,_,_,_,_=await self.task_fixture(source_id)
        target=(await self.client.get('/api/customers/'+str(self.customer_id))).json()
        source=(await self.client.get('/api/customers/'+str(source_id))).json()
        merged=await self.client.post(f'/api/customers/{source_id}/merge',json={'target_id':self.customer_id,
            'expected_revision':source['revision'],'expected_target_revision':target['revision'],
            'reason':'核对同一主体资料合并','identity_checked':True,'dry_run':False})
        self.assertEqual(merged.status_code,200,merged.text)
        await self.login('clerk')
        url='/api/billing-tasks/'+merged_task['id']
        self.assertEqual((await self.client.post(url+'/reveal',json={'purpose':'合并后查看'})).status_code,409)
        self.assertEqual((await self.client.get(url)).json()['status'],'needs_review')

if __name__ == "__main__":
    unittest.main()
