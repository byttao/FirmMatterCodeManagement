"""Persistent FIFO exports, private artifacts and current-permission download gates."""
import asyncio
from datetime import datetime, timedelta
import hashlib
import json
import os
from pathlib import Path
import secrets
import shutil
import threading
import time
from typing import Literal, Optional
from fastapi import APIRouter, Depends, Header, HTTPException, Query, Request
from fastapi.responses import StreamingResponse
from pydantic import Field
from sqlalchemy import text
from sqlalchemy.orm import joinedload, selectinload
from openpyxl import Workbook
from auth import get_current_user
from database import DATA_DIR, SessionLocal, get_db
from permissions import has, project_scope, can_read_money
from schemas import InputModel
from project_queries import filter_projects
from runtime_log import event as log_event
from audit import event, write_lock
import models
from limited_response import LimitedStreamingResponse

DIRECTORY=DATA_DIR/'exports'
MAX_ROWS=int(os.getenv('FIRM_EXPORT_MAX_ROWS','50000'))
MAX_SECONDS=int(os.getenv('FIRM_EXPORT_MAX_SECONDS','300'))
MAX_BYTES=int(os.getenv('FIRM_EXPORT_MAX_BYTES',str(50*1024*1024)))
MAX_PENDING=int(os.getenv('FIRM_EXPORT_MAX_PENDING','20'))
DOWNLOAD_RATE=max(1024,int(os.getenv('FIRM_EXPORT_DOWNLOAD_KIB','100'))*1024)
download_slot=threading.BoundedSemaphore(1)
router=APIRouter(prefix='/api/export-jobs')
PROJECT_FIELDS={'project_id':'项目ID','customer_name':'客户名称','customer_tax_id':'税号',
 'firm':'事务所','report_type':'报告类型','report_year':'业务年度','report_no':'报告编号',
 'report_status':'编号状态','leader':'负责人','project_status':'项目状态','contract_amount':'合同金额',
 'invoiced_amount':'开票金额','received_amount':'收款金额','created_at':'创建时间'}
MONEY={'contract_amount','invoiced_amount','received_amount'}


class ExportFilters(InputModel):
    fiscal_year: int = Field(...,ge=2000,le=2100)
    search: Optional[str] = Field(None,max_length=200)
    firm: Optional[str] = Field(None,max_length=100)
    report_type: Optional[str] = Field(None,max_length=100)
    project_status: Optional[str] = Field(None,max_length=20)
    leader_id: Optional[int] = Field(None,ge=1)
    report_year: Optional[int] = Field(None,ge=2000,le=2100)
    customer_id: Optional[int] = Field(None,ge=1)


class ExportCreate(InputModel):
    export_type: Literal['projects','signers','billing_sensitive']='projects'
    filters: Optional[ExportFilters]=None
    columns: list[str] = Field(default_factory=list,max_items=30)
    purpose: Optional[str] = Field(None,max_length=500)
    confirm_sensitive: bool=False


def permission(user,kind):
    if kind=='projects':
        return any(has(user,p) for p in ('project.export.all','project.export.finance','project.export.related'))
    if kind=='signers':return has(user,'signer.manage')
    return has(user,'billing.export_sensitive') and has(user,'billing.read.all')


def epoch(db):
    return db.get(models.ExportAccessEpoch,1).value


def valid_owner(db,job):
    user=db.get(models.User,job.requested_by)
    return user and user.is_active and permission(user,job.export_type) and user.permission_revision==job.permission_revision and epoch(db)==job.access_epoch


def view(job,db):
    position=None
    if job.status=='queued':
        position=db.query(models.ExportJob.id).filter(models.ExportJob.status=='queued',models.ExportJob.created_at<=job.created_at).count()
    return {k:getattr(job,k) for k in ('id','export_type','status','created_at','started_at','finished_at',
        'expires_at','row_count','file_size','error_code','request_id')} | {'queue_position':position}


@router.post('',status_code=202)
def submit(data:ExportCreate,request:Request,key:str=Header(...,alias='Idempotency-Key',min_length=8,max_length=100),
           user=Depends(get_current_user),db=Depends(get_db)):
    if not permission(user,data.export_type):raise HTTPException(403,'无权限导出此类资料')
    if data.export_type=='billing_sensitive' and (not data.confirm_sensitive or not data.purpose or not data.purpose.strip()):
        raise HTTPException(422,'敏感批量导出须确认并填写用途')
    if data.export_type=='projects':
        if not data.filters or not db.query(models.FiscalYear.id).filter_by(year=data.filters.fiscal_year).first():
            raise HTTPException(422,'请指定已配置的编号年度')
        columns=data.columns or [k for k in PROJECT_FIELDS if k not in MONEY or has(user,'finance.read.all') or has(user,'finance.summary.led')]
        if any(k not in PROJECT_FIELDS for k in columns) or len(set(columns))!=len(columns):raise HTTPException(422,'导出列无效')
        if any(k in MONEY for k in columns) and not (has(user,'finance.read.all') or has(user,'finance.summary.led')):
            raise HTTPException(403,'无权限导出金额列')
    else:
        if data.filters is not None:raise HTTPException(422,'此类导出为全范围资料，不接受项目筛选')
        if data.columns:raise HTTPException(422,'此类导出使用固定资料列')
        columns=[]
    values=data.dict();digest=hashlib.sha256(json.dumps(values,sort_keys=True,ensure_ascii=False).encode()).hexdigest()
    db.rollback();write_lock(db)
    prior=db.query(models.ExportJob).filter_by(requested_by=user.id,idempotency_key=key).first()
    if prior:
        if prior.payload_hash!=digest:raise HTTPException(409,'同幂等键内容不一致')
        return view(prior,db)
    pending=db.query(models.ExportJob).filter(models.ExportJob.status.in_(['queued','running']))
    if pending.count()>=MAX_PENDING or pending.filter_by(requested_by=user.id).count()>=2:
        raise HTTPException(429,'导出队列已满，请等待现有任务完成')
    job=models.ExportJob(id=secrets.token_hex(16),requested_by=user.id,export_type=data.export_type,
        filters_json=json.dumps(values['filters'] or {},ensure_ascii=False),field_set=json.dumps(columns),purpose=data.purpose,
        permission_revision=user.permission_revision,access_epoch=epoch(db),idempotency_key=key,payload_hash=digest,
        request_id=request.state.request_id)
    db.add(job);db.flush();event(db,user,'export.submit',job,reason=data.purpose,request=request)
    db.commit();return view(job,db)


def get_job(db,job_id,user,manage=False):
    job=db.get(models.ExportJob,job_id)
    if not job or (job.requested_by!=user.id and not (manage and has(user,'identity.manage'))):raise HTTPException(404,'任务不存在')
    return job


@router.get('')
def listing(page:int=Query(1,ge=1),manage:bool=False,user=Depends(get_current_user),db=Depends(get_db)):
    query=db.query(models.ExportJob)
    if manage and not has(user,'identity.manage'):raise HTTPException(403,'无权限管理队列')
    if not manage:query=query.filter_by(requested_by=user.id)
    return {'items':[view(j,db) for j in query.order_by(models.ExportJob.created_at.desc(),models.ExportJob.id).offset((page-1)*20).limit(20)],'total':query.count(),'page':page}


@router.get('/{job_id}')
def detail(job_id:str,user=Depends(get_current_user),db=Depends(get_db)):
    return view(get_job(db,job_id,user,True),db)


@router.post('/{job_id}/cancel')
def cancel(job_id:str,request:Request,user=Depends(get_current_user),db=Depends(get_db)):
    write_lock(db);job=get_job(db,job_id,user,True)
    if job.status not in ('queued','running'):raise HTTPException(409,'任务已结束')
    job.cancel_requested=True
    if job.status=='queued':job.status='cancelled';job.finished_at=datetime.utcnow()
    event(db,user,'export.cancel',job,request=request);db.commit();return view(job,db)


@router.get('/{job_id}/download')
def download(job_id:str,request:Request,user=Depends(get_current_user),db=Depends(get_db)):
    job=get_job(db,job_id,user,True)
    if not permission(user,job.export_type) or not valid_owner(db,job):raise HTTPException(403,'资料权限已变动，请重新提交导出')
    # An administrator may download another user's artifact only with full scope.
    if job.requested_by!=user.id and job.export_type=='projects' and not has(user,'project.export.all'):
        raise HTTPException(403,'无权限下载他人的资料范围')
    if job.status=='expired' or job.expires_at and datetime.utcnow()>=job.expires_at:raise HTTPException(410,'文件已过期，请重新导出')
    if job.status!='succeeded':raise HTTPException(409,'文件尚未生成完成')
    path=DIRECTORY/(job.file_id+'.xlsx')
    if not path.is_file():raise HTTPException(410,'文件已清理，请重新导出')
    if not download_slot.acquire(blocking=False):raise HTTPException(429,'已有文件正在下载，请稍后重试',headers={'Retry-After':'5'})
    try:
        event(db,user,'export.download',job,request=request);db.commit()
    except BaseException:
        download_slot.release();raise
    async def stream():
        started=time.monotonic();sent=0
        try:
            with path.open('rb') as source:
                while True:
                    chunk=await asyncio.to_thread(source.read,16384)
                    if not chunk:break
                    sent+=len(chunk)
                    await asyncio.sleep(max(0,sent/DOWNLOAD_RATE-(time.monotonic()-started)))
                    yield chunk
        finally:
            log_event('export.transfer.finished',job_id=job_id,duration_ms=round((time.monotonic()-started)*1000),result='complete' if sent==job.file_size else 'interrupted')
    return LimitedStreamingResponse(stream(),slot=download_slot,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition':f'attachment; filename="{job.export_type}_{job.id[:8]}.xlsx"',
                 'Content-Length':str(job.file_size),'Cache-Control':'no-store'})


class ExportStopped(Exception):
    pass


def check(job_id,started):
    if time.monotonic()-started>MAX_SECONDS:raise ExportStopped('budget_exceeded')
    wal=DATA_DIR/'db.sqlite-wal'
    if wal.exists() and wal.stat().st_size>256*1024*1024:raise ExportStopped('wal_budget_exceeded')
    with SessionLocal() as db:
        job=db.get(models.ExportJob,job_id)
        if job.cancel_requested:raise ExportStopped('cancelled')
        if not valid_owner(db,job):raise ExportStopped('permission_changed')


def text_cell(value):
    if isinstance(value,str) and value.lstrip().startswith(('=','+','-','@')):return "'"+value
    return value


def generate(job_id):
    DIRECTORY.mkdir(parents=True,exist_ok=True)
    temporary=DIRECTORY/(job_id+'.partial.xlsx');ready=DIRECTORY/(job_id+'.xlsx')
    started=time.monotonic();count=0;book=None
    try:
        if shutil.disk_usage(DIRECTORY).free<100*1024*1024:raise ExportStopped('disk_space_low')
        check(job_id,started)
        with SessionLocal() as snapshot:
            snapshot.execute(text('BEGIN'))  # One real WAL read snapshot for all batches.
            job=snapshot.get(models.ExportJob,job_id);user=snapshot.get(models.User,job.requested_by)
            filters=json.loads(job.filters_json);columns=json.loads(job.field_set)
            book=Workbook(write_only=True)
            info=book.create_sheet('导出说明')
            for row in [('资料类型',job.export_type),('筛选条件',json.dumps(filters,ensure_ascii=False)),
                        ('数据读取时间UTC',datetime.utcnow().isoformat()+'Z'),('用途',job.purpose or ''),
                        ('范围说明','按生成时权限与筛选导出；无权限金额留空')]:info.append([text_cell(v) for v in row])
            sheet=book.create_sheet('资料')
            if job.export_type=='projects':
                query=filter_projects(project_scope(snapshot.query(models.Project),user,exporting=True),filters)
                sheet.append([PROJECT_FIELDS[k] for k in columns])
                query=query.options(joinedload(models.Project.leader),selectinload(models.Project.members))
                def values(p):
                    result={k:getattr(p,k,None) for k in columns if k not in ('leader','report_status')}
                    result.update(leader=p.leader.real_name if p.leader else '',report_status={'pending':'待编号','assigned':'已编号','recycled':'已作废'}.get(p.report_no_status,p.report_no_status))
                    for k in MONEY:
                        if not can_read_money(user,p):result[k]=None
                    return [result[k].isoformat()+'Z' if isinstance(result[k],datetime) else result[k] for k in columns]
                key=models.Project.id
            elif job.export_type=='signers':
                query=snapshot.query(models.Signer).options(joinedload(models.Signer.user));key=models.Signer.id
                sheet.append(['姓名','执业账号','事务所','状态'])
                def values(s):return [s.name,s.user.username if s.user else '',s.signer_type,'启用' if s.is_active else '禁用']
            else:
                query=snapshot.query(models.BillingProfile).filter_by(is_active=True);key=models.BillingProfile.id
                sheet.append(['客户名称','税号','档案','银行账号','开户行','地址','电话','联系人','邮箱'])
                def values(profile):
                    from data_crypto import decrypt
                    version=snapshot.get(models.BillingVersion,profile.current_verified_version_id)
                    if not version:return None
                    sensitive=decrypt(version.sensitive_ciphertext,version.key_id,f'billing-version:{version.id}')
                    customer=snapshot.get(models.Customer,profile.customer_id)
                    return [customer.name,customer.tax_id,profile.label]+[sensitive.get(k,'') for k in ('bank_account','bank_name','registered_address','registered_phone','contact_name','recipient_email')]
            last=0
            while True:
                check(job_id,started)
                batch=query.filter(key>last).order_by(key).limit(500).all()
                if not batch:break
                for item in batch:
                    row=values(item)
                    if row is not None:sheet.append([text_cell(v) for v in row]);count+=1
                    if count>MAX_ROWS:raise ExportStopped('row_limit_exceeded')
                last=batch[-1].id
                snapshot.expunge_all()
                with SessionLocal() as progress:
                    progress.query(models.ExportJob).filter_by(id=job_id,status='running').update({'row_count':count});progress.commit()
            book.save(temporary);book.close();book=None
        check(job_id,started)
        if temporary.stat().st_size>MAX_BYTES:raise ExportStopped('file_limit_exceeded')
        temporary.chmod(0o600);temporary.replace(ready)
        with SessionLocal() as db:
            write_lock(db);job=db.get(models.ExportJob,job_id)
            if job.cancel_requested or not valid_owner(db,job):raise ExportStopped('permission_changed')
            job.status='succeeded';job.file_id=job.id;job.file_size=ready.stat().st_size;job.row_count=count
            job.finished_at=datetime.utcnow();job.expires_at=job.finished_at+timedelta(hours=24)
            db.commit()
        log_event('export.finished',job_id=job_id,result='succeeded',row_count=count)
    except Exception as error:
        code=str(error) if isinstance(error,ExportStopped) else 'generation_failed'
        temporary.unlink(missing_ok=True);ready.unlink(missing_ok=True)
        with SessionLocal() as db:
            job=db.get(models.ExportJob,job_id)
            job.status='cancelled' if code=='cancelled' else 'failed';job.error_code=code;job.finished_at=datetime.utcnow();db.commit()
        log_event('export.finished',job_id=job_id,result='failed',reason_code=code)
    finally:
        if book:
            for sheet in book.worksheets:
                try:sheet.close();sheet._writer.cleanup()
                except Exception:pass
            book.close()


def cleanup(restart=False):
    DIRECTORY.mkdir(parents=True,exist_ok=True)
    with SessionLocal() as db:
        if restart:
            db.query(models.ExportJob).filter_by(status='running').update({'status':'failed','error_code':'interrupted','finished_at':datetime.utcnow()})
            for path in DIRECTORY.glob('*.partial.xlsx'):path.unlink(missing_ok=True)
        for job in db.query(models.ExportJob).filter(models.ExportJob.expires_at<=datetime.utcnow(),models.ExportJob.status=='succeeded'):
            (DIRECTORY/(job.file_id+'.xlsx')).unlink(missing_ok=True);job.status='expired'
        db.query(models.ExportJob).filter(models.ExportJob.finished_at<datetime.utcnow()-timedelta(days=30)).delete()
        db.commit()


def claim():
    with SessionLocal() as db:
        write_lock(db)
        if db.query(models.ExportJob.id).filter_by(status='running').first():return None
        job=db.query(models.ExportJob).filter_by(status='queued').order_by(models.ExportJob.created_at,models.ExportJob.id).first()
        if not job:return None
        job.status='running';job.started_at=datetime.utcnow();db.commit();return job.id
