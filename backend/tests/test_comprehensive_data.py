"""End-to-end data-flow coverage for the current API contract.

The test deliberately uses a temporary data directory.  It covers the positive
path of every business API family and the cross-numbered-year rule: ``fiscal_year``
is the numbering year while ``report_year`` is the audited business year.
"""

import importlib
import os
from pathlib import Path
import sys
import tempfile
import unittest
from io import BytesIO

import httpx
from openpyxl import Workbook, load_workbook


class ComprehensiveDataTest(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.backend_dir = Path(__file__).resolve().parents[1]
        self.temp = tempfile.TemporaryDirectory()
        self.old_cwd = Path.cwd()
        self.old_data_dir = os.environ.get("FIRM_MANAGER_DATA_DIR")
        self.old_database_url = os.environ.get("FIRM_MANAGER_DATABASE_URL")
        os.environ["FIRM_MANAGER_DATA_DIR"] = self.temp.name
        os.environ.pop("FIRM_MANAGER_DATABASE_URL", None)
        sys.path.insert(0, str(self.backend_dir))
        os.chdir(self.backend_dir)
        # Other integration tests import the application with another database.
        # Remove those modules so this test always gets a fresh engine.
        for name in ("main", "database", "models", "schemas", "migrations", "finance", "auth", "report_no_generator"):
            sys.modules.pop(name, None)
        self.main = importlib.import_module("main")
        self.client = httpx.AsyncClient(
            transport=httpx.ASGITransport(app=self.main.app, client=("127.0.0.1", 23456)),
            base_url="http://testserver",
        )

    async def asyncTearDown(self):
        await self.client.aclose()
        self.main.engine.dispose()
        os.chdir(self.old_cwd)
        if self.old_data_dir is None:
            os.environ.pop("FIRM_MANAGER_DATA_DIR", None)
        else:
            os.environ["FIRM_MANAGER_DATA_DIR"] = self.old_data_dir
        if self.old_database_url is None:
            os.environ.pop("FIRM_MANAGER_DATABASE_URL", None)
        else:
            os.environ["FIRM_MANAGER_DATABASE_URL"] = self.old_database_url
        sys.path.remove(str(self.backend_dir))
        for name in ("main", "database", "models", "schemas", "migrations", "finance", "auth", "report_no_generator"):
            sys.modules.pop(name, None)
        self.temp.cleanup()

    async def login(self, username, password, fiscal_year=None):
        payload = {"username": username, "password": password}
        if fiscal_year is not None:
            payload["fiscal_year"] = fiscal_year
        response = await self.client.post("/api/auth/login", json=payload)
        self.assertEqual(response.status_code, 200, response.text)
        return {"Authorization": f"Bearer {response.json()['access_token']}"}

    async def test_all_business_flows_and_cross_numbered_years(self):
        setup = {
            "admin_username": "office_admin",
            "admin_password": "a-strong-test-password",
            "admin_real_name": "管理员",
            "fiscal_year": 2026,
            "firms": [
                {"name": "甲会计师事务所", "report_types": [
                    {"report_type": "年度审计", "template": "AF-{yyyy}-{nnn}"},
                    {"report_type": "专项审计", "template": "SP-{yyyy}-{yy}-{nn}"},
                ]},
                {"name": "乙税务师事务所", "report_types": [
                    {"report_type": "税务鉴证", "template": "TX-{yyyy}-{nnn}"},
                ]},
            ],
        }
        self.assertEqual((await self.client.get("/api/setup/status")).json(), {"initialized": False})
        self.assertEqual((await self.client.get("/api/public/numbered-years")).json(), {"years": []})
        self.assertEqual((await self.client.post("/api/auth/login", json={
            "username": "office_admin", "password": setup["admin_password"],
        })).status_code, 428)
        setup_response = await self.client.post("/api/setup", json=setup)
        self.assertEqual(setup_response.status_code, 200, setup_response.text)
        self.assertEqual((await self.client.get("/api/public/numbered-years")).json(), {"years": [2026]})
        self.assertEqual((await self.client.post("/api/setup", json=setup)).status_code, 409)

        admin = await self.login("office_admin", setup["admin_password"])
        self.assertEqual((await self.client.get("/api/auth/me", headers=admin)).json()["role"], "admin")
        self.assertEqual((await self.client.post("/api/auth/login", json={
            "username": "office_admin", "password": setup["admin_password"], "fiscal_year": 2027,
        })).status_code, 400)

        options = await self.client.get("/api/numbered-years/options", headers=admin, params={
            "year": 2026, "firm": "甲会计师事务所",
        })
        self.assertEqual(options.status_code, 200, options.text)
        self.assertEqual(set(options.json()["report_types"]), {"年度审计", "专项审计"})
        self.assertEqual((await self.client.get("/api/numbered-years/options", headers=admin)).status_code, 200)
        config = await self.client.get("/api/numbered-years", headers=admin)
        self.assertEqual(config.status_code, 200, config.text)
        first_year = next(item for item in config.json() if item["year"] == 2026)
        first_firm = next(item for item in first_year["firms"] if item["firm"] == "甲会计师事务所")
        audit_type = next(item for item in first_firm["report_types"] if item["report_type"] == "年度审计")

        # User lifecycle, duplicate names, search and role restrictions.
        users = {}
        for username, role in (("auditor_one", "practitioner"), ("auditor_two", "practitioner"), ("office_staff", "admin_staff")):
            response = await self.client.post("/api/users", headers=admin, json={
                "username": username, "password": "another-password", "real_name": "同名执业人" if role == "practitioner" else "行政人员", "role": role,
            })
            self.assertEqual(response.status_code, 200, response.text)
            users[username] = response.json()
        self.assertEqual((await self.client.get("/api/users", headers=admin)).status_code, 200)
        self.assertEqual(len((await self.client.get("/api/users/practitioners", headers=admin)).json()), 2)
        self.assertEqual({row["username"] for row in (await self.client.get("/api/users/search", headers=admin, params={"q": "tongming"})).json()}, {"auditor_one", "auditor_two"})
        self.assertEqual({row["username"] for row in (await self.client.get("/api/users/search", headers=admin, params={"q": "同名"})).json()}, {"auditor_one", "auditor_two"})
        staff_auth = await self.login("office_staff", "another-password")
        practitioner_auth = await self.login("auditor_one", "another-password")
        self.assertEqual((await self.client.get("/api/users", headers=staff_auth)).status_code, 403)

        # Signer lifecycle and duplicate-name disambiguation by user account.
        signer_one = await self.client.post("/api/signers", headers=admin, json={"user_id": users["auditor_one"]["id"], "signer_type": "甲会计师事务所"})
        signer_two = await self.client.post("/api/signers", headers=admin, json={"user_id": users["auditor_two"]["id"], "signer_type": "甲会计师事务所"})
        self.assertEqual(signer_one.status_code, 200, signer_one.text)
        self.assertEqual(signer_two.status_code, 200, signer_two.text)
        self.assertEqual(signer_one.json()["name"], signer_two.json()["name"])
        self.assertNotEqual(signer_one.json()["user"]["username"], signer_two.json()["user"]["username"])
        self.assertEqual((await self.client.post("/api/signers", headers=admin, json={"user_id": users["auditor_one"]["id"], "signer_type": "甲会计师事务所"})).status_code, 400)
        renamed = await self.client.put(f"/api/users/{users['auditor_two']['id']}", headers=admin, json={"real_name": "同名执业人（二）"})
        self.assertEqual(renamed.status_code, 200, renamed.text)
        refreshed_signer_two = (await self.client.get("/api/signers", headers=admin)).json()
        self.assertEqual(next(item for item in refreshed_signer_two if item["id"] == signer_two.json()["id"])["name"], "同名执业人（二）")
        self.assertEqual((await self.client.get("/api/signers", headers=admin, params={"eligible_only": True})).status_code, 200)
        self.assertEqual((await self.client.get("/api/signers/template", headers=admin)).status_code, 200)
        self.assertEqual((await self.client.get("/api/signers/export", headers=admin)).status_code, 200)
        self.assertEqual((await self.client.post("/api/signers/999999/disable", headers=admin)).status_code, 404)
        self.assertEqual((await self.client.post(f"/api/signers/{signer_two.json()['id']}/disable", headers=admin)).status_code, 200)
        self.assertEqual((await self.client.post(f"/api/signers/{signer_two.json()['id']}/enable", headers=admin)).status_code, 200)
        self.assertEqual((await self.client.put(f"/api/signers/{signer_two.json()['id']}", headers=admin, json={"signer_type": "乙税务师事务所"})).status_code, 200)
        self.assertEqual((await self.client.put(f"/api/signers/{signer_two.json()['id']}", headers=admin, json={"signer_type": "甲会计师事务所"})).status_code, 200)

        # Two business years under one numbering year share the same sequence.
        common = {"firm": "甲会计师事务所", "report_type": "年度审计", "customer_name": "A公司", "leader_id": users["auditor_one"]["id"], "signer1_id": signer_one.json()["id"], "contract_amount": 1000}
        project_2024 = await self.client.post("/api/projects", headers=admin, json={**common, "report_year": 2024, "member_ids": [users["auditor_two"]["id"]]})
        project_2025 = await self.client.post("/api/projects", headers=admin, json={**common, "report_year": 2025, "customer_name": "B公司", "signer1_id": signer_two.json()["id"]})
        self.assertEqual(project_2024.status_code, 200, project_2024.text)
        self.assertEqual(project_2025.status_code, 200, project_2025.text)
        p24, p25 = project_2024.json(), project_2025.json()
        self.assertEqual(p24["fiscal_year"], p25["fiscal_year"])
        self.assertEqual(p24["project_id"], "PRJ-2026-0001")
        self.assertEqual(p25["project_id"], "PRJ-2026-0002")
        self.assertEqual((await self.client.post(f"/api/projects/{p24['project_id']}/generate-report-no", headers=admin)).json()["report_no"], "AF-2026-001")
        self.assertEqual((await self.client.post(f"/api/projects/{p25['project_id']}/generate-report-no", headers=admin)).json()["report_no"], "AF-2026-002")
        self.assertEqual((await self.client.get("/api/projects", headers=admin, params={"report_year": 2024})).json()["total"], 1)
        self.assertEqual((await self.client.get("/api/projects", headers=admin, params={"search": "B公司", "firm": "甲会计师事务所", "report_type": "年度审计", "leader_id": users["auditor_one"]["id"]})).json()["total"], 1)
        self.assertEqual((await self.client.get("/api/projects", headers=admin, params={"page": 1, "page_size": 1})).json()["page_size"], 1)
        self.assertEqual((await self.client.get("/api/projects", headers=admin, params={"year": 2024})).status_code, 400)
        self.assertEqual((await self.client.get(f"/api/projects/{p24['id']}", headers=admin)).status_code, 200)
        self.assertEqual((await self.client.get(f"/api/projects/{p24['project_id']}", headers=admin)).status_code, 200)
        self.assertEqual((await self.client.get(f"/api/projects/{p24['project_id']}/report-number-history", headers=admin)).status_code, 200)

        # Finance ledger: each invoice/receipt is independent and totals refresh.
        invoice_url = f"/api/projects/{p24['id']}/finance/invoices"
        receipt_url = f"/api/projects/{p24['id']}/finance/receipts"
        invoice = await self.client.post(invoice_url, headers=staff_auth, json={"amount": "100.00", "occurred_on": "2026-01-10", "reference": "INV-1"})
        invoice_two = await self.client.post(invoice_url, headers=admin, json={"amount": "50.55", "occurred_on": "2026-02-10", "reference": "INV-2"})
        receipt = await self.client.post(receipt_url, headers=staff_auth, json={"amount": "40.00", "occurred_on": "2026-03-01", "reference": "REC-1"})
        self.assertEqual(invoice.status_code, 200, invoice.text)
        self.assertEqual(invoice_two.status_code, 200, invoice_two.text)
        self.assertEqual(receipt.status_code, 200, receipt.text)
        updated_invoice = await self.client.put(f"{invoice_url}/{invoice.json()['id']}", headers=staff_auth, json={"amount": "110.00"})
        self.assertEqual(updated_invoice.status_code, 200, updated_invoice.text)
        current_project = (await self.client.get(f"/api/projects/{p24['id']}", headers=admin)).json()
        self.assertEqual(current_project["invoiced_amount"], 160.55)
        self.assertEqual(current_project["received_amount"], 40.0)
        self.assertEqual(current_project["uninvoiced_amount"], 839.45)
        self.assertEqual((await self.client.get(invoice_url, headers=practitioner_auth)).status_code, 200)
        self.assertEqual((await self.client.delete(f"{receipt_url}/{receipt.json()['id']}", headers=staff_auth)).status_code, 200)
        self.assertEqual((await self.client.post(invoice_url, headers=practitioner_auth, json={"amount": 1, "occurred_on": "2026-04-01"})).status_code, 403)
        self.assertEqual((await self.client.delete(f"/api/users/{users['office_staff']['id']}", headers=admin)).status_code, 200)
        disabled_users = await self.client.get("/api/users", headers=admin, params={"include_disabled": True})
        self.assertEqual(disabled_users.status_code, 200)
        self.assertFalse(next(item for item in disabled_users.json() if item["username"] == "office_staff")["is_active"])

        # A signer's view follows the logged-in numbering year, while report_year stays independent.
        self.assertEqual((await self.client.get("/api/projects/signed-by-me", headers=practitioner_auth)).json()["total"], 1)
        fy_2025 = await self.client.post("/api/numbered-years", headers=admin, json={"year": 2025})
        fy_2025_id = fy_2025.json()["id"]
        firm_2025 = await self.client.post("/api/numbered-years/firms", headers=admin, json={"fiscal_year_id": fy_2025_id, "firm": "甲会计师事务所"})
        rule_2025 = await self.client.post("/api/numbered-years/rules", headers=admin, json={"fiscal_year_firm_id": firm_2025.json()["id"], "rule_name": "年度审计编号", "template": "AF-{yyyy}-{nnn}"})
        type_2025 = await self.client.post("/api/numbered-years/report-types", headers=admin, json={"fiscal_year_firm_id": firm_2025.json()["id"], "report_type": "年度审计", "rule_id": rule_2025.json()["id"]})
        self.assertEqual(type_2025.status_code, 200, type_2025.text)
        self.assertEqual((await self.client.put("/api/users/current-fiscal-year", headers=admin, json={"fiscal_year": 2025})).status_code, 200)
        practitioner_2025 = await self.login("auditor_one", "another-password", 2025)
        p2025 = await self.client.post("/api/projects", headers=practitioner_2025, json={**common, "report_year": 2024, "customer_name": "C公司"})
        self.assertEqual(p2025.status_code, 200, p2025.text)
        self.assertEqual((await self.client.post(f"/api/projects/{p2025.json()['project_id']}/generate-report-no", headers=practitioner_2025)).json()["report_no"], "AF-2025-001")
        self.assertEqual([p["customer_name"] for p in (await self.client.get("/api/projects/signed-by-me", headers=practitioner_2025)).json()["items"]], ["C公司"])
        practitioner_2026 = await self.login("auditor_one", "another-password", 2026)
        self.assertEqual([p["customer_name"] for p in (await self.client.get("/api/projects/signed-by-me", headers=practitioner_2026)).json()["items"]], ["A公司"])
        dashboard = await self.client.get("/api/dashboard", headers=admin, params={"fiscal_year": 2026})
        self.assertEqual(dashboard.status_code, 200, dashboard.text)
        self.assertEqual(dashboard.json()["total_projects"], 2)
        self.assertEqual((await self.client.get("/api/dashboard", headers=admin, params={"fiscal_year": 2025})).json()["total_projects"], 1)
        exported = await self.client.get("/api/export/projects", headers=admin, params={"fiscal_year": 2026})
        self.assertEqual(exported.status_code, 200, exported.text)
        sheet = load_workbook(BytesIO(exported.content), data_only=False).active
        self.assertEqual({row[0].value for row in sheet.iter_rows(min_row=2)}, {"PRJ-2026-0001", "PRJ-2026-0002"})
        self.assertEqual((await self.client.get("/api/export/projects", headers=practitioner_auth)).status_code, 403)

        # Number placeholder overflow: {nn} renders 100 as 100, without truncating it.
        switch_back = await self.client.put("/api/users/current-fiscal-year", headers=admin, json={"fiscal_year": 2026})
        self.assertEqual(switch_back.status_code, 200, switch_back.text)
        admin = {"Authorization": f"Bearer {switch_back.json()['access_token']}"}
        short_rule = await self.client.post("/api/numbered-years/rules", headers=admin, json={"fiscal_year_firm_id": first_firm["id"], "rule_name": "两位序号", "template": "SN-{yyyy}-{nn}"})
        self.assertEqual(short_rule.status_code, 200, short_rule.text)
        short_type = await self.client.post("/api/numbered-years/report-types", headers=admin, json={"fiscal_year_firm_id": first_firm["id"], "report_type": "短序号业务", "rule_id": short_rule.json()["id"]})
        self.assertEqual(short_type.status_code, 200, short_type.text)
        with self.main.SessionLocal() as db:
            db.query(self.main.models.ReportNumberRule).filter_by(id=short_rule.json()["id"]).one().current_sequence = 99
            db.commit()
        short_project = await self.client.post("/api/projects", headers=admin, json={**common, "report_type": "短序号业务", "report_year": 2026, "customer_name": "序号溢出测试", "signer1_id": None})
        self.assertEqual(short_project.status_code, 200, short_project.text)
        self.assertEqual((await self.client.post(f"/api/projects/{short_project.json()['project_id']}/generate-report-no", headers=admin)).json()["report_no"], "SN-2026-100")

        # Recycle/delete and configuration CRUD.
        self.assertEqual((await self.client.post(f"/api/projects/{p25['project_id']}/recycle-report-no", headers=admin)).status_code, 200)
        self.assertEqual((await self.client.post(f"/api/projects/{p25['project_id']}/recycle-report-no", headers=admin)).status_code, 400)
        self.assertEqual((await self.client.delete(f"/api/projects/{short_project.json()['project_id']}", headers=admin)).status_code, 200)
        self.assertEqual((await self.client.get(f"/api/projects/{short_project.json()['project_id']}", headers=admin)).status_code, 404)
        self.assertEqual((await self.client.delete(f"/api/numbered-years/report-types/{short_type.json()['id']}", headers=admin)).status_code, 400)

        transient_year = await self.client.post("/api/numbered-years", headers=admin, json={"year": 2027})
        transient_firm = await self.client.post("/api/numbered-years/firms", headers=admin, json={"fiscal_year_id": transient_year.json()["id"], "firm": "临时事务所"})
        transient_rule = await self.client.post("/api/numbered-years/rules", headers=admin, json={"fiscal_year_firm_id": transient_firm.json()["id"], "rule_name": "临时规则", "template": "TMP-{yyyy}-{n}"})
        transient_type = await self.client.post("/api/numbered-years/report-types", headers=admin, json={"fiscal_year_firm_id": transient_firm.json()["id"], "report_type": "临时业务", "rule_id": transient_rule.json()["id"]})
        self.assertEqual((await self.client.put(f"/api/numbered-years/rules/{transient_rule.json()['id']}", headers=admin, json={"rule_name": "改名规则"})).status_code, 200)
        self.assertEqual((await self.client.put(f"/api/numbered-years/report-types/{transient_type.json()['id']}", headers=admin, json={"report_type": "改名业务"})).status_code, 200)
        self.assertEqual((await self.client.delete(f"/api/numbered-years/report-types/{transient_type.json()['id']}", headers=admin)).status_code, 200)
        self.assertEqual((await self.client.delete(f"/api/numbered-years/rules/{transient_rule.json()['id']}", headers=admin)).status_code, 200)
        self.assertEqual((await self.client.delete(f"/api/numbered-years/firms/{transient_firm.json()['id']}", headers=admin)).status_code, 200)
        self.assertEqual((await self.client.delete(f"/api/numbered-years/{transient_year.json()['id']}", headers=admin)).status_code, 200)

        # Import path is tested after all positive CRUD paths, using a new practitioner.
        imported_user = await self.client.post("/api/users", headers=admin, json={"username": "auditor_three", "password": "another-password", "real_name": "导入签字人", "role": "practitioner"})
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["执业账号", "事务所"])
        sheet.append(["auditor_three", "甲会计师事务所"])
        sheet.append(["auditor_three", "甲会计师事务所"])
        uploaded = await self.client.post("/api/signers/import", headers=admin, files={"file": ("signers.xlsx", BytesIO(self._workbook_bytes(workbook)), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        self.assertEqual(uploaded.status_code, 200, uploaded.text)
        self.assertEqual(uploaded.json()["errors"][0].find("重复") >= 0, True)
        self.assertEqual((await self.client.post("/api/signers/999999/enable", headers=admin)).status_code, 404)
        imported_signer = next(item for item in (await self.client.get("/api/signers", headers=admin)).json() if item["user_id"] == imported_user.json()["id"])
        self.assertEqual((await self.client.delete(f"/api/signers/{imported_signer['id']}", headers=admin)).status_code, 200)

    @staticmethod
    def _workbook_bytes(workbook):
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()


if __name__ == "__main__":
    unittest.main()
