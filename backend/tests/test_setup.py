"""Integration checks against a fresh, isolated database."""

import importlib
import asyncio
import os
from pathlib import Path
import sys
import tempfile
import unittest
from io import BytesIO

import httpx
from openpyxl import Workbook, load_workbook


class FirstRunTest(unittest.IsolatedAsyncioTestCase):
    def test_signer_account_migration_preserves_legacy_rows(self):
        from sqlalchemy import create_engine, inspect, text
        backend_dir = Path(__file__).resolve().parents[1]
        sys.path.insert(0, str(backend_dir))
        try:
            from migrations import migrate_signer_accounts
            with tempfile.TemporaryDirectory() as data_dir:
                engine = create_engine(f"sqlite:///{Path(data_dir) / 'old.sqlite'}")
                with engine.begin() as connection:
                    connection.execute(text("CREATE TABLE users (id INTEGER PRIMARY KEY)"))
                    connection.execute(text(
                        "CREATE TABLE signers (id INTEGER PRIMARY KEY, name TEXT, signer_type TEXT)"
                    ))
                    connection.execute(text(
                        "INSERT INTO signers (id, name, signer_type) VALUES (1, 'Old Signer', 'Old Firm')"
                    ))
                migrate_signer_accounts(engine)
                migrate_signer_accounts(engine)
                self.assertIn("user_id", {column["name"] for column in inspect(engine).get_columns("signers")})
                with engine.connect() as connection:
                    self.assertEqual(connection.execute(text("SELECT name, user_id FROM signers")).one(),
                                     ("Old Signer", None))
                engine.dispose()
        finally:
            sys.path.remove(str(backend_dir))

    async def test_setup_and_number_lifecycle(self):
        backend_dir = Path(__file__).resolve().parents[1]
        with tempfile.TemporaryDirectory() as data_dir:
            old_cwd = Path.cwd()
            old_data_dir = os.environ.get("FIRM_MANAGER_DATA_DIR")
            os.environ["FIRM_MANAGER_DATA_DIR"] = data_dir
            sys.path.insert(0, str(backend_dir))
            os.chdir(backend_dir)
            try:
                app = importlib.import_module("main").app
                local_transport = httpx.ASGITransport(app=app, client=("127.0.0.1", 12345))
                async with httpx.AsyncClient(transport=local_transport, base_url="http://testserver") as client:
                    status = await client.get("/api/setup/status")
                    self.assertEqual(status.json(), {"initialized": False})

                    setup = {
                        "admin_username": "office_admin",
                        "admin_password": "a-strong-test-password",
                        "admin_real_name": "管理员",
                        "fiscal_year": 2026,
                        "firms": [{
                            "name": "Example Firm",
                            "report_types": [{"report_type": "Audit", "template": "EX-{yyyy}-{nnn}"}],
                        }],
                    }
                    invalid = await client.post("/api/setup", json={
                        **setup,
                        "firms": [{"name": "Example Firm", "report_types": [{"report_type": "Audit", "template": "EX-{yyyy}"}]}],
                    })
                    self.assertEqual(invalid.status_code, 400)
                    self.assertEqual((await client.get("/api/setup/status")).json(), {"initialized": False})
                    duplicate_formats = await client.post("/api/setup", json={
                        **setup,
                        "firms": [
                            setup["firms"][0],
                            {"name": "Another Firm", "report_types": [
                                {"report_type": "Other Audit", "template": "EX-{yyyy}-{nnnn}"},
                            ]},
                        ],
                    })
                    self.assertEqual(duplicate_formats.status_code, 400)

                    remote_transport = httpx.ASGITransport(app=app, client=("192.0.2.1", 12345))
                    async with httpx.AsyncClient(transport=remote_transport, base_url="http://testserver") as remote:
                        self.assertEqual((await remote.post("/api/setup", json=setup)).status_code, 403)

                    self.assertEqual((await client.post("/api/setup", json=setup)).status_code, 200)
                    self.assertEqual((await client.get("/api/setup/status")).json(), {"initialized": True})
                    self.assertEqual((await client.post("/api/setup", json=setup)).status_code, 409)

                    login = await client.post("/api/auth/login", json={
                        "username": setup["admin_username"], "password": setup["admin_password"],
                    })
                    self.assertEqual(login.status_code, 200, login.text)
                    auth = {"Authorization": f"Bearer {login.json()['access_token']}"}
                    years = await client.get("/api/numbered-years", headers=auth)
                    self.assertEqual(years.json()[0]["firms"][0]["rules"][0]["template"], "EX-{yyyy}-{nnn}")
                    configured_year = years.json()[0]
                    first_firm = configured_year["firms"][0]
                    first_rule_id = first_firm["rules"][0]["id"]
                    first_type_id = first_firm["report_types"][0]["id"]

                    self.assertEqual((await client.put("/api/users/current-fiscal-year", headers=auth, json={
                        "fiscal_year": 2027,
                    })).status_code, 400)
                    second_firm = await client.post("/api/numbered-years/firms", headers=auth, json={
                        "fiscal_year_id": configured_year["id"], "firm": "Another Firm",
                    })
                    self.assertEqual(second_firm.status_code, 200, second_firm.text)
                    second_firm_id = second_firm.json()["id"]
                    self.assertEqual((await client.post("/api/numbered-years/rules", headers=auth, json={
                        "fiscal_year_firm_id": second_firm_id, "rule_name": "Duplicate",
                        "template": "EX-{yyyy}-{nnnn}",
                    })).status_code, 400)
                    invalid_rule = await client.post("/api/numbered-years/rules", headers=auth, json={
                        "fiscal_year_firm_id": second_firm_id, "rule_name": "Invalid",
                        "template": "AF-{yyyy}-{bad}",
                    })
                    self.assertEqual(invalid_rule.status_code, 400)
                    second_rule = await client.post("/api/numbered-years/rules", headers=auth, json={
                        "fiscal_year_firm_id": second_firm_id, "rule_name": "Standard",
                        "template": "AF-{yyyy}-{nnnn}",
                    })
                    self.assertEqual(second_rule.status_code, 200, second_rule.text)
                    self.assertEqual(second_rule.json()["sequence_digits"], 4)
                    second_rule_id = second_rule.json()["id"]
                    self.assertEqual((await client.put(f"/api/numbered-years/rules/{second_rule_id}", headers=auth, json={
                        "template": "AF-{yyyy}-{nnn}", "sequence_digits": 4,
                    })).status_code, 400)
                    self.assertEqual((await client.put(f"/api/numbered-years/rules/{second_rule_id}", headers=auth, json={
                        "template": "EX-{yyyy}-{nnn}", "sequence_digits": 3,
                    })).status_code, 400)
                    self.assertEqual((await client.post("/api/numbered-years/report-types", headers=auth, json={
                        "fiscal_year_firm_id": second_firm_id, "report_type": "Other Audit",
                        "rule_id": first_rule_id,
                    })).status_code, 400)
                    second_type = await client.post("/api/numbered-years/report-types", headers=auth, json={
                        "fiscal_year_firm_id": second_firm_id, "report_type": "Other Audit",
                        "rule_id": second_rule_id,
                    })
                    self.assertEqual(second_type.status_code, 200, second_type.text)
                    self.assertEqual((await client.put(f"/api/numbered-years/report-types/{first_type_id}", headers=auth, json={
                        "rule_id": second_rule_id,
                    })).status_code, 400)
                    self.assertEqual((await client.delete(f"/api/numbered-years/rules/{first_rule_id}", headers=auth)).status_code, 400)

                    empty_year = await client.post("/api/numbered-years", headers=auth, json={"year": 2025})
                    self.assertEqual(empty_year.status_code, 200, empty_year.text)
                    empty_firm = await client.post("/api/numbered-years/firms", headers=auth, json={
                        "fiscal_year_id": empty_year.json()["id"], "firm": "Temporary Firm",
                    })
                    empty_rule = await client.post("/api/numbered-years/rules", headers=auth, json={
                        "fiscal_year_firm_id": empty_firm.json()["id"], "rule_name": "Temporary",
                        "template": "TMP-{yyyy}-{nnn}",
                    })
                    self.assertEqual((await client.post("/api/numbered-years/report-types", headers=auth, json={
                        "fiscal_year_firm_id": empty_firm.json()["id"], "report_type": "Temporary",
                        "rule_id": empty_rule.json()["id"],
                    })).status_code, 200)
                    self.assertEqual((await client.delete(f"/api/numbered-years/{empty_year.json()['id']}", headers=auth)).status_code, 200)
                    self.assertFalse(any(year["year"] == 2025 for year in (await client.get("/api/numbered-years", headers=auth)).json()))

                    practitioner = await client.post("/api/users", headers=auth, json={
                        "username": "auditor", "password": "another-password", "real_name": "执业人员", "role": "practitioner",
                    })
                    self.assertEqual(practitioner.status_code, 200, practitioner.text)
                    project = await client.post("/api/projects", headers=auth, json={
                        "firm": "Example Firm", "report_type": "Audit", "report_year": 2026,
                        "customer_name": "Sample Client", "leader_id": practitioner.json()["id"],
                    })
                    self.assertEqual(project.status_code, 200, project.text)
                    project_id = project.json()["project_id"]
                    self.assertEqual((await client.put(f"/api/numbered-years/report-types/{first_type_id}", headers=auth, json={
                        "report_type": "Renamed Audit",
                    })).status_code, 400)
                    self.assertEqual((await client.delete(f"/api/numbered-years/report-types/{first_type_id}", headers=auth)).status_code, 400)
                    self.assertEqual((await client.delete(f"/api/numbered-years/firms/{first_firm['id']}", headers=auth)).status_code, 400)
                    self.assertEqual((await client.delete(f"/api/numbered-years/{configured_year['id']}", headers=auth)).status_code, 400)

                    other_user = await client.post("/api/users", headers=auth, json={
                        "username": "other_auditor", "password": "another-password", "real_name": "其他人员", "role": "practitioner",
                    })
                    self.assertEqual(other_user.status_code, 200)
                    other_login = await client.post("/api/auth/login", json={
                        "username": "other_auditor", "password": "another-password",
                    })
                    other_auth = {"Authorization": f"Bearer {other_login.json()['access_token']}"}
                    forbidden = await client.post(f"/api/projects/{project_id}/generate-report-no", headers=other_auth)
                    self.assertEqual(forbidden.status_code, 403)

                    first = await client.post(f"/api/projects/{project_id}/generate-report-no", headers=auth)
                    self.assertEqual(first.status_code, 200, first.text)
                    self.assertEqual(first.json()["report_no"], "EX-2026-001")
                    recycled = await client.post(f"/api/projects/{project_id}/recycle-report-no", headers=auth)
                    self.assertEqual(recycled.status_code, 200, recycled.text)
                    self.assertEqual((await client.get(f"/api/projects/{project_id}", headers=auth)).json()["report_no_status"], "recycled")
                    self.assertEqual((await client.post(f"/api/projects/{project_id}/recycle-report-no", headers=auth)).status_code, 400)
                    second = await client.post(f"/api/projects/{project_id}/generate-report-no", headers=auth)
                    self.assertEqual(second.status_code, 200, second.text)
                    self.assertEqual(second.json()["report_no"], "EX-2026-002")

                    # Simulate an upgraded database with undated aggregate balances.
                    main = importlib.import_module("main")
                    models = importlib.import_module("models")
                    finance = importlib.import_module("finance")
                    with main.SessionLocal() as db:
                        stored = db.query(models.Project).filter_by(project_id=project_id).one()
                        stored.invoiced_amount = 100.02
                        stored.received_amount = 30.01
                        db.query(models.FinanceMigration).filter_by(project_id=stored.id).delete()
                        db.commit()
                    finance.migrate_legacy_finance()
                    finance.migrate_legacy_finance()
                    invoices_url = f"/api/projects/{project_id}/finance/invoices"
                    receipts_url = f"/api/projects/{project_id}/finance/receipts"
                    legacy_invoice = (await client.get(invoices_url, headers=auth)).json()
                    self.assertEqual(len(legacy_invoice), 1)
                    self.assertTrue(legacy_invoice[0]["is_legacy"])
                    self.assertIsNone(legacy_invoice[0]["occurred_on"])
                    self.assertEqual(legacy_invoice[0]["amount"], 100.02)
                    self.assertEqual(len((await client.get(receipts_url, headers=auth)).json()), 1)

                    self.assertEqual((await client.post(invoices_url, headers=auth, json={
                        "amount": 0, "occurred_on": "2026-09-24",
                    })).status_code, 422)
                    invoice = await client.post(invoices_url, headers=auth, json={
                        "amount": 49.98, "occurred_on": "2026-09-24", "reference": "INV-01",
                    })
                    self.assertEqual(invoice.status_code, 200, invoice.text)
                    receipt = await client.post(receipts_url, headers=auth, json={
                        "amount": 20, "occurred_on": "2026-09-25",
                    })
                    self.assertEqual(receipt.status_code, 200, receipt.text)
                    current = (await client.get(f"/api/projects/{project_id}", headers=auth)).json()
                    self.assertEqual(current["invoiced_amount"], 150)
                    self.assertEqual(current["received_amount"], 50.01)
                    self.assertEqual(current["unreceived_amount"], 99.99)
                    self.assertTrue(current["invoice_date"].startswith("2026-09-24"))

                    edited = await client.put(f"{invoices_url}/{invoice.json()['id']}", headers=auth, json={
                        "amount": 40, "occurred_on": "2026-09-26",
                    })
                    self.assertEqual(edited.status_code, 200, edited.text)
                    self.assertEqual((await client.get(f"/api/projects/{project_id}", headers=auth)).json()["invoiced_amount"], 140.02)
                    deleted = await client.delete(f"{receipts_url}/{receipt.json()['id']}", headers=auth)
                    self.assertEqual(deleted.status_code, 200)
                    self.assertEqual((await client.get(f"/api/projects/{project_id}", headers=auth)).json()["received_amount"], 30.01)

                    self.assertEqual((await client.put(f"/api/projects/{project_id}", headers=other_auth, json={
                        "member_ids": [other_user.json()["id"]],
                    })).status_code, 403)
                    staff = await client.post("/api/users", headers=auth, json={
                        "username": "office_staff", "password": "another-password",
                        "real_name": "行政人员", "role": "admin_staff",
                    })
                    staff_login = await client.post("/api/auth/login", json={
                        "username": "office_staff", "password": "another-password",
                    })
                    staff_auth = {"Authorization": f"Bearer {staff_login.json()['access_token']}"}
                    self.assertEqual((await client.put(f"/api/projects/{project_id}", headers=staff_auth, json={
                        "member_ids": [staff.json()["id"]],
                    })).status_code, 403)
                    self.assertEqual((await client.put(f"/api/projects/{project_id}", headers=auth, json={
                        "report_year": 2027,
                    })).status_code, 400)
                    self.assertEqual((await client.put(f"/api/projects/{project_id}", headers=auth, json={
                        "invoiced_amount": 900,
                    })).status_code, 422)
                    self.assertEqual((await client.get(invoices_url, headers=other_auth)).status_code, 403)
                    self.assertEqual((await client.post(invoices_url, headers=other_auth, json={
                        "amount": 10, "occurred_on": "2026-09-24",
                    })).status_code, 403)

                    new_project = {
                        "firm": "Example Firm", "report_type": "Audit", "report_year": 2026,
                        "customer_name": "Sample Client", "leader_id": practitioner.json()["id"],
                    }
                    created = await asyncio.gather(*[
                        client.post("/api/projects", headers=auth, json=new_project) for _ in range(2)
                    ])
                    self.assertTrue(all(result.status_code == 200 for result in created), [r.text for r in created])
                    self.assertEqual(len({r.json()["project_id"] for r in created}), 2)
                    numbered = await asyncio.gather(*[
                        client.post(f"/api/projects/{r.json()['project_id']}/generate-report-no", headers=auth)
                        for r in created
                    ])
                    self.assertTrue(all(result.status_code == 200 for result in numbered), [r.text for r in numbered])
                    self.assertEqual(len({r.json()["report_no"] for r in numbered}), 2)

                    # Distinct accounts with the same legal name remain distinct signers.
                    self.assertEqual((await client.put(
                        f"/api/users/{other_user.json()['id']}", headers=auth, json={"real_name": "执业人员"}
                    )).status_code, 200)
                    first_signer = await client.post("/api/signers", headers=auth, json={
                        "user_id": practitioner.json()["id"], "signer_type": "Example Firm",
                    })
                    second_signer = await client.post("/api/signers", headers=auth, json={
                        "user_id": other_user.json()["id"], "signer_type": "Example Firm",
                    })
                    self.assertEqual(first_signer.status_code, 200, first_signer.text)
                    self.assertEqual(second_signer.status_code, 200, second_signer.text)
                    self.assertEqual(first_signer.json()["name"], second_signer.json()["name"])
                    self.assertNotEqual(first_signer.json()["user"]["username"], second_signer.json()["user"]["username"])
                    signer_export = await client.get("/api/signers/export", headers=auth)
                    self.assertEqual(signer_export.status_code, 200, signer_export.text)
                    signer_sheet = load_workbook(BytesIO(signer_export.content), data_only=False).active
                    self.assertEqual({row[1] for row in signer_sheet.iter_rows(min_row=2, values_only=True)},
                                     {"auditor", "other_auditor"})
                    self.assertEqual((await client.post("/api/signers", headers=auth, json={
                        "user_id": staff.json()["id"], "signer_type": "Example Firm",
                    })).status_code, 400)
                    self.assertEqual((await client.post("/api/signers", headers=auth, json={
                        "user_id": other_user.json()["id"], "signer_type": "Example Firm",
                    })).status_code, 400)

                    signed = await client.post("/api/projects", headers=auth, json={
                        **new_project, "customer_name": "=1+1",
                        "signer1_id": second_signer.json()["id"],
                    })
                    self.assertEqual(signed.status_code, 200, signed.text)
                    signed_id = signed.json()["project_id"]
                    self.assertEqual((await client.get(f"/api/projects/{signed_id}", headers=other_auth)).status_code, 200)
                    mine = await client.get("/api/projects/signed-by-me", headers=other_auth)
                    self.assertEqual(mine.status_code, 200, mine.text)
                    self.assertEqual([item["project_id"] for item in mine.json()["items"]], [signed_id])
                    dashboard = await client.get("/api/dashboard", headers=other_auth)
                    self.assertEqual(dashboard.status_code, 200, dashboard.text)
                    self.assertEqual(dashboard.json()["total_projects"], 1)
                    exported = await client.get("/api/export/projects", headers=auth)
                    self.assertEqual(exported.status_code, 200, exported.text)
                    sheet = load_workbook(BytesIO(exported.content), data_only=False).active
                    columns = {cell.value: index for index, cell in enumerate(sheet[1])}
                    exported_row = next(row for row in sheet.iter_rows(min_row=2) if row[0].value == signed_id)
                    self.assertEqual(exported_row[columns["签字人一账号"]].value, "other_auditor")
                    self.assertEqual(exported_row[columns["客户名称"]].value, "=1+1")
                    self.assertEqual(exported_row[columns["客户名称"]].data_type, "s")
                    self.assertEqual((await client.get("/api/projects/signed-by-me", headers=staff_auth)).status_code, 403)
                    self.assertIn(signed_id, [item["project_id"] for item in (
                        await client.get("/api/projects", headers=other_auth)
                    ).json()["items"]])
                    self.assertEqual((await client.post("/api/projects", headers=auth, json={
                        **new_project, "signer1_id": 999999,
                    })).status_code, 400)
                    self.assertEqual((await client.post("/api/projects", headers=auth, json={
                        **new_project, "leader_id": staff.json()["id"],
                    })).status_code, 400)

                    # Old signer references survive upgrade until an administrator links an account.
                    with main.SessionLocal() as db:
                        legacy = models.Signer(name="旧签字人", signer_type="Example Firm", is_active=True)
                        db.add(legacy)
                        db.commit()
                        legacy_id = legacy.id
                        old_project = db.query(models.Project).filter_by(project_id=project_id).one()
                        old_project.signer1_id = legacy_id
                        db.commit()
                    self.assertNotIn(legacy_id, [item["id"] for item in (
                        await client.get("/api/signers/by-firm/Example Firm", headers=auth)
                    ).json()])
                    self.assertEqual((await client.put(f"/api/projects/{project_id}", headers=auth, json={
                        "project_phase": "归档",
                    })).status_code, 200)
                    third = await client.post("/api/users", headers=auth, json={
                        "username": "third_auditor", "password": "another-password",
                        "real_name": "执业人员", "role": "practitioner",
                    })
                    third_login = await client.post("/api/auth/login", json={
                        "username": "third_auditor", "password": "another-password",
                    })
                    third_auth = {"Authorization": f"Bearer {third_login.json()['access_token']}"}
                    self.assertEqual((await client.get(f"/api/projects/{project_id}", headers=third_auth)).status_code, 403)
                    linked = await client.put(f"/api/signers/{legacy_id}", headers=auth, json={
                        "user_id": third.json()["id"],
                    })
                    self.assertEqual(linked.status_code, 200, linked.text)
                    self.assertEqual(linked.json()["name"], "旧签字人")
                    self.assertEqual((await client.get(f"/api/projects/{project_id}", headers=third_auth)).status_code, 200)
                    self.assertEqual((await client.put(f"/api/signers/{legacy_id}", headers=auth, json={
                        "user_id": other_user.json()["id"],
                    })).status_code, 400)
                    fourth = await client.post("/api/users", headers=auth, json={
                        "username": "fourth_auditor", "password": "another-password",
                        "real_name": "执业人员", "role": "practitioner",
                    })
                    workbook = Workbook()
                    workbook.active.append(["执业账号", "事务所"])
                    workbook.active.append(["fourth_auditor", "Example Firm"])
                    workbook.active.append(["office_staff", "Example Firm"])
                    output = BytesIO()
                    workbook.save(output)
                    imported = await client.post("/api/signers/import", headers=auth, files={
                        "file": ("signers.xlsx", output.getvalue(),
                                 "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
                    })
                    self.assertEqual(imported.status_code, 200, imported.text)
                    self.assertEqual(imported.json()["message"], "导入完成：新增 1 条，跳过 1 条")
                    self.assertTrue(any(item["user_id"] == fourth.json()["id"] for item in (
                        await client.get("/api/signers/by-firm/Example Firm", headers=auth)
                    ).json()))
                    self.assertEqual((await client.put(f"/api/users/{third.json()['id']}", headers=auth, json={
                        "role": "admin_staff",
                    })).status_code, 200)
                    self.assertNotIn(legacy_id, [item["id"] for item in (
                        await client.get("/api/signers/by-firm/Example Firm", headers=auth)
                    ).json()])
                    self.assertEqual((await client.post("/api/projects", headers=auth, json={
                        **new_project, "signer1_id": legacy_id,
                    })).status_code, 400)
                    self.assertEqual((await client.put(f"/api/projects/{project_id}", headers=auth, json={
                        "project_phase": "归档",
                    })).status_code, 200)
                    cleared = await client.put(f"/api/projects/{project_id}", headers=auth, json={
                        "signer1_id": None,
                    })
                    self.assertEqual(cleared.status_code, 200, cleared.text)
                    self.assertIsNone(cleared.json()["signer1_id"])
                    admin_id = (await client.get("/api/auth/me", headers=auth)).json()["id"]
                    self.assertEqual((await client.put(f"/api/users/{admin_id}", headers=auth, json={
                        "role": "practitioner",
                    })).status_code, 400)
                    self.assertEqual((await client.put(f"/api/users/{admin_id}", headers=auth, json={
                        "is_active": False,
                    })).status_code, 400)
            finally:
                os.chdir(old_cwd)
                sys.path.remove(str(backend_dir))
                if old_data_dir is None:
                    os.environ.pop("FIRM_MANAGER_DATA_DIR", None)
                else:
                    os.environ["FIRM_MANAGER_DATA_DIR"] = old_data_dir


if __name__ == "__main__":
    unittest.main()
